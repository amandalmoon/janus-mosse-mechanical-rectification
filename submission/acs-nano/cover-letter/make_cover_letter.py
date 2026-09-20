from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pathlib import Path

out=Path('/mnt/data/janus_n19_acs_submission/cover_letter/ACS_Nano_Cover_Letter_DRAFT.docx')
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(0.7); sec.bottom_margin=Inches(0.7); sec.left_margin=Inches(0.9); sec.right_margin=Inches(0.9)
st=doc.styles['Normal']; st.font.name='Liberation Serif'; st._element.rPr.rFonts.set(qn('w:ascii'),'Liberation Serif'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Liberation Serif'); st.font.size=Pt(10.5)

def para(text='',after=5,first=0):
    p=doc.add_paragraph(text); p.paragraph_format.space_after=Pt(after)
    if first: p.paragraph_format.first_line_indent=Inches(first)
    return p

p=para('[DATE]',4); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; p.runs[0].bold=True
para('[CORRESPONDING AUTHOR FULL NAME]',0)
para('[AFFILIATION]',0)
para('[FULL MAILING ADDRESS]',0)
para('[EMAIL]',8)
para('ACS Nano',0); para('Editorial Office',8)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(8); p.add_run('Re: ').bold=True; p.add_run('Submission of an Article, “Unguided Full-2D Mechanical Rectification and Vector Mode Locking in Finite Janus MoSSe Contacts”')
para('Dear Editor:',7)
para('Please consider our manuscript, “Unguided Full-2D Mechanical Rectification and Vector Mode Locking in Finite Janus MoSSe Contacts,” for publication as an Article in ACS Nano. The coauthors, in manuscript order, are [FULL COAUTHOR NAMES].',5,0.22)
para('The manuscript presents a DFT-anchored computational study of finite Janus MoSSe contacts in the full two-dimensional registry plane. Starting from the published MoSSe generalized stacking-fault-energy landscape, we resolve prepared-state directional depinning, test the mechanism with matched same-spectrum symmetry controls, examine compact-contact scaling and boundary-registry competition, and determine zero-mean lattice-vector transport with Floquet, stationarity, thermal, and in-plane-relaxation robustness checks. The claims are explicitly conditioned on state preparation, loading axis, and the declared reduced dynamical protocol rather than presented as parameter-free experimental-device predictions.',5,0.22)
para('We believe the work fits ACS Nano because it connects atomically resolved interfacial energetics in a Janus two-dimensional material with finite-contact nanomechanics, surface/interface physics, and nonlinear nanoscale transport using a reproducible theory-and-simulation framework.',5,0.22)
para('The Supporting Information contains the prepared-state branch audit; matched symmetry controls; registry-origin diagnostics; extended size/shape and boundary-sensitivity tests; drive-period/Floquet checks; stationary thermal statistics; linear and nonlinear in-plane relaxation cross-checks; and reproducibility details. A separate reproducibility package is prepared for submission or repository deposition as appropriate.',5,0.22)
para('Preprint status: [NO PREPRINT / PREPRINT SERVER AND DOI OR URL]. Prior discussion with an ACS Nano editor: [NONE / DETAILS]. Review-Only Material beyond the Supporting Information: [NONE / DESCRIPTION]. The authors will confirm originality, exclusive submission, authorship approval, funding, and conflict-of-interest information in the submission system before upload.',5,0.22)
para('Thank you for your consideration.',7)
para('Sincerely,',2)
para('[CORRESPONDING AUTHOR NAME]',0)
para('on behalf of all authors',0)
para('[INSTITUTION]',0)
para('[EMAIL]',0)
doc.save(out)
print(out)
