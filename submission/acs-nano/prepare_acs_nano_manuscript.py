from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from copy import deepcopy
from pathlib import Path
import re, json

SRC=Path('/mnt/data/janus_n18_publication_figures/Janus_MoSSe_PublicationFigure_Manuscript_v7.docx')
OUT=Path('/mnt/data/janus_n19_acs_submission/ACS_Nano_Initial_Submission_Manuscript.docx')

def ptext(el, doc):
    if el.tag == qn('w:p'):
        return Paragraph(el, doc._body).text.strip()
    return None

def remove_paragraph(p):
    el=p._element
    el.getparent().remove(el)
    p._p=p._element=None

def insert_paragraph_after(paragraph, text='', style=None):
    new_p = OxmlElement('w:p')
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    if style:
        new_para.style=style
    return new_para

def expand_cite_group(s):
    nums=[]
    for token in re.split(r'\s*,\s*', s):
        token=token.strip()
        m=re.fullmatch(r'(\d+)\s*-\s*(\d+)', token)
        if m:
            a,b=map(int,m.groups())
            nums.extend(range(a,b+1) if b>=a else [a,b])
        elif token.isdigit():
            nums.append(int(token))
    return nums

def compress_nums(nums):
    nums=sorted(dict.fromkeys(nums))
    if not nums: return ''
    out=[]; i=0
    while i<len(nums):
        j=i
        while j+1<len(nums) and nums[j+1]==nums[j]+1:
            j+=1
        if j-i>=2:
            out.append(f'{nums[i]}–{nums[j]}')
        elif j-i==1:
            out.extend([str(nums[i]),str(nums[j])])
        else:
            out.append(str(nums[i]))
        i=j+1
    return ','.join(out)

CITE_RE=re.compile(r'\[(\d+(?:\s*(?:-|,)\s*\d+)*)\]')

def iter_blocks(doc):
    body=doc._element.body
    for child in body.iterchildren():
        if child.tag==qn('w:p'):
            yield Paragraph(child, doc._body)
        elif child.tag==qn('w:tbl'):
            yield Table(child, doc._body)

def texts_in_block(block):
    if isinstance(block, Paragraph):
        yield block
    else:
        for row in block.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p

def collect_citation_order(doc):
    order=[]; seen=set()
    for block in iter_blocks(doc):
        if isinstance(block, Paragraph) and block.text.strip()=='References':
            break
        for p in texts_in_block(block):
            for m in CITE_RE.finditer(p.text):
                for n in expand_cite_group(m.group(1)):
                    if n not in seen:
                        order.append(n); seen.add(n)
    return order

def clear_runs(p):
    for r in list(p.runs):
        p._p.remove(r._r)

def rewrite_citations(p, mapping):
    text=p.text
    if not CITE_RE.search(text): return
    # preserve whole-paragraph direct formatting if it was uniform
    base = p.runs[0] if p.runs else None
    bold = base.bold if base else None
    italic = base.italic if base else None
    font_name = base.font.name if base else None
    font_size = base.font.size if base else None
    clear_runs(p)
    pos=0
    for m in CITE_RE.finditer(text):
        pre=text[pos:m.start()]
        nextpos=m.end()
        punct=''
        if nextpos < len(text) and text[nextpos] in '.,;:':
            punct=text[nextpos]
            nextpos += 1
        r=p.add_run(pre + punct)
        if base:
            r.bold=bold; r.italic=italic
            if font_name: r.font.name=font_name
            if font_size: r.font.size=font_size
        oldnums=expand_cite_group(m.group(1))
        newnums=[mapping[n] for n in oldnums]
        cr=p.add_run(compress_nums(newnums))
        cr.font.superscript=True
        if font_name: cr.font.name=font_name
        if font_size: cr.font.size=font_size
        pos=nextpos
    r=p.add_run(text[pos:])
    if base:
        r.bold=bold; r.italic=italic
        if font_name: r.font.name=font_name
        if font_size: r.font.size=font_size

def italicize_phrase(p, phrase):
    # after citation rewrite, apply simple phrase-level splitting across individual runs
    for run in list(p.runs):
        if run.font.superscript: continue
        txt=run.text
        idx=txt.lower().find(phrase.lower())
        if idx<0: continue
        # split current run with XML insertion
        parent=run._r.getparent(); at=parent.index(run._r)
        before=txt[:idx]; mid=txt[idx:idx+len(phrase)]; after=txt[idx+len(phrase):]
        props=deepcopy(run._r.rPr) if run._r.rPr is not None else None
        parent.remove(run._r)
        def mk(text, ital=False):
            rr=OxmlElement('w:r')
            if props is not None: rr.append(deepcopy(props))
            t=OxmlElement('w:t'); t.text=text
            if text.startswith(' ') or text.endswith(' '): t.set(qn('xml:space'),'preserve')
            rr.append(t)
            if ital:
                rpr=rr.get_or_add_rPr(); i=OxmlElement('w:i'); rpr.append(i)
            return rr
        new=[]
        if before: new.append(mk(before))
        new.append(mk(mid,True))
        if after: new.append(mk(after))
        for rr in reversed(new): parent.insert(at,rr)
        break

doc=Document(SRC)
body=doc._element.body

# Remove working subtitle and unheaded Introduction label
for p in list(doc.paragraphs):
    t=p.text.strip()
    if t.startswith('A DFT-anchored finite-contact model with prepared-state'):
        remove_paragraph(p)
        break
for p in list(doc.paragraphs):
    if p.text.strip()=='1. Introduction':
        remove_paragraph(p)
        break

# Move the complete Methods block (including equations/tables) to the end of main text, before Data Availability
children=list(body)
start=next(i for i,e in enumerate(children) if ptext(e,doc)=='2. Model and computational methods')
end=next(i for i,e in enumerate(children) if ptext(e,doc)=='3. Results')
method_block=children[start:end]
for e in method_block: body.remove(e)
children=list(body)
data_anchor=next(e for e in children if ptext(e,doc)=='6. Data and code availability')
for e in method_block:
    data_anchor.addprevious(e)

# Unnumber headings and conform key names
special={
    'Model and computational methods':'Methods',
    'Conclusion':'Conclusions',
    'Data and code availability':'Data Availability Statement',
}
for p in doc.paragraphs:
    if p.style.name.startswith('Heading'):
        t=re.sub(r'^\d+(?:\.\d+)*\.?\s*','',p.text.strip())
        t=special.get(t,t)
        p.text=t

# Keywords: 7 lowercase terms; chemical formula moved out to avoid forced lowercase distortion.
for p in doc.paragraphs:
    if p.text.strip().startswith('Keywords:'):
        p.text=('Keywords: janus transition-metal dichalcogenides; mechanical rectification; depinning; '
                'finite-contact scaling; structural lubricity; vector mode locking; rocking ratchet')
        break

# Title treatment
p0=doc.paragraphs[0]
p0.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in p0.runs:
    r.bold=True; r.font.size=Pt(16)

# Add substantial AI-use disclosure in Methods, after reproducibility paragraph
for p in doc.paragraphs:
    if p.text.strip().startswith('The published-GSFE engine, analytic gradient and Hessian'):
        ai=insert_paragraph_after(p,
            'OpenAI ChatGPT was used as an AI-assisted research support tool for code drafting and debugging, '
            'audit orchestration, manuscript language revision, and figure-layout preparation. Numerical claims '
            'were accepted only when produced by executable scripts and checked by clean reruns, cross-solver '
            'calculations, or independent implementations as appropriate; responsibility for the scientific '
            'interpretation and submitted content remains with the authors.',
            style='Body Text')
        break

# Add current-guideline back matter before References.
ref_heading=next(p for p in doc.paragraphs if p.text.strip()=='References')
si_h=ref_heading.insert_paragraph_before('Supporting Information')
si_h.style='Heading 1'
si_p=ref_heading.insert_paragraph_before(
    'Supporting Information: Prepared-state branch enumeration and continuation; matched symmetry and exact-inversion controls; '
    'translation-minimized inversion-asymmetry diagnostics; extended compact-contact size/shape similarity; boundary-registry '
    'competition and edge-weight sensitivity; drive-period and Floquet checks; stationary thermal trajectory statistics; '
    'linear finite-element and nonlinear valence-force relaxation checks; and reproducibility details (PDF).')
si_p.style='Body Text'
ack_h=ref_heading.insert_paragraph_before('Acknowledgments')
ack_h.style='Heading 1'
ack_p=ref_heading.insert_paragraph_before(
    'OpenAI ChatGPT was used during manuscript preparation for AI-assisted code drafting and debugging, workflow organization, '
    'language revision, and figure-layout preparation. The authors are responsible for all scientific content and submitted material.')
ack_p.style='Body Text'

# Complete the three references that previously used et al. despite <20 authors.
ref_replacements={
6: 'Liao, M.; Nicolini, P.; Du, L.; Yuan, J.; Wang, S.; Yu, H.; Tang, J.; Cheng, P.; Watanabe, K.; Taniguchi, T.; Gu, L.; Claerbout, V. E. P.; Silva, A.; Kramer, D.; Polcar, T.; Yang, R.; Shi, D.; Zhang, G. Ultra-low friction and edge-pinning effect in large-lattice-mismatch van der Waals heterostructures. Nature Materials 21, 47-53 (2022). https://doi.org/10.1038/s41563-021-01058-4.',
9: 'Lu, A.-Y.; Zhu, H.; Xiao, J.; Chuu, C.-P.; Han, Y.; Chiu, M.-H.; Cheng, C.-C.; Yang, C.-W.; Wei, K.-H.; Yang, Y.; Wang, Y.; Sokaras, D.; Nordlund, D.; Yang, P.; Muller, D. A.; Chou, M.-Y.; Zhang, X.; Li, L.-J. Janus monolayers of transition metal dichalcogenides. Nature Nanotechnology 12, 744-749 (2017). https://doi.org/10.1038/nnano.2017.100.',
10:'Zhang, K.; Guo, Y.; Ji, Q.; Lu, A.-Y.; Su, C.; Wang, H.; Puretzky, A. A.; Geohegan, D.; Qian, X.; Fang, S.; Kaxiras, E.; Kong, J.; Huang, S. Enhancement of van der Waals interlayer coupling through polar Janus MoSSe. Journal of the American Chemical Society 142, 17499-17507 (2020). https://doi.org/10.1021/jacs.0c07051.'
}
# capture old references
refs={}
ref_paras=[]
seen_ref=False
for p in doc.paragraphs:
    if p.text.strip()=='References':
        seen_ref=True; continue
    if seen_ref:
        m=re.match(r'^(\d+)\.\s+(.*)$',p.text.strip())
        if m:
            n=int(m.group(1)); refs[n]=ref_replacements.get(n,m.group(2)); ref_paras.append(p)

# Determine first-citation order in the reordered manuscript and create new numbering.
order=collect_citation_order(doc)
# include any uncited references at the end as a warning-preserving fallback (should not occur)
for n in sorted(refs):
    if n not in order: order.append(n)
mapping={old:i+1 for i,old in enumerate(order)}

# Convert in-text bracket citations to ACS-style superscript numbers and renumber them.
for block in iter_blocks(doc):
    if isinstance(block, Paragraph) and block.text.strip()=='References':
        break
    for p in texts_in_block(block):
        rewrite_citations(p,mapping)
        if 'et al.' in p.text:
            italicize_phrase(p,'et al.')

# Rewrite reference paragraphs in new citation order; Fast Format allows any complete consistent reference style.
for i,p in enumerate(ref_paras):
    old=order[i]
    p.text=f'{i+1}. {refs[old]}'

# Consistent unnumbered heading formatting
for p in doc.paragraphs:
    if p.style.name=='Heading 1':
        for r in p.runs:
            r.bold=True
    if p.text.strip()=='Abstract':
        for r in p.runs: r.bold=True

# Save mapping provenance
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
(Path('/mnt/data/janus_n19_acs_submission/qa/citation_renumber_map.json')).write_text(json.dumps(mapping,indent=2),encoding='utf-8')
print('saved',OUT)
print('citation order old->new',mapping)
