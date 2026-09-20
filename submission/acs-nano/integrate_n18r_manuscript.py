from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches

REPLACEMENTS = {
    "cycle-resolved integer lattice-vector transport over all 91 sampled amplitude-period points":
        "cycle-resolved integer winding states over all 91 sampled amplitude-period points, including pinned and transporting states",
    "all 91 sampled combinations in the F0*=1.0-4.0 and tau*=20-80 grid reproduce integer lattice transport":
        "all 91 sampled combinations in the F0*=1.0-4.0 and tau*=20-80 grid reproduce cycle-resolved integer winding states, including pinned and transporting states",
    "cycle-resolved integer lattice-vector transport over all 91 sampled amplitude-period points, with representative plateaus independently verified as attracting":
        "cycle-resolved integer winding states over all 91 sampled amplitude-period points, including pinned and transporting states, with representative transporting plateaus independently verified as attracting",
}

FIGURE_STEMS = {
    1: "Figure_1_symmetry_contracts_N18R.png",
    2: "Figure_2_branch_resolved_depinning_N18R.png",
    3: "Figure_3_extended_size_shape_similarity_N18R.png",
    4: "Figure_4_boundary_registry_switching_N18R.png",
    5: "Figure_5_vector_mode_map_floquet_N18R.png",
    6: "Figure_6_thermal_stationary_hierarchy_N18R.png",
}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--manuscript", type=Path, required=True)
    ap.add_argument("--figure-dir", type=Path, required=True)
    ap.add_argument("--manifest-dir", type=Path, required=True)
    ap.add_argument("--output", type=Path, required=True)
    ap.add_argument("--embedded-width-in", type=float, default=6.40)
    args = ap.parse_args()

    doc = Document(args.manuscript)

    replacement_hits = {old: 0 for old in REPLACEMENTS}
    for p in doc.paragraphs:
        for run in p.runs:
            for old, new in REPLACEMENTS.items():
                if old in run.text:
                    replacement_hits[old] += run.text.count(old)
                    run.text = run.text.replace(old, new)

    if any(v != 1 for v in replacement_hits.values()):
        raise AssertionError(f"unexpected wording-replacement counts: {replacement_hits}")

    captions: dict[int, str] = {}
    for i in range(1, 7):
        manifest = json.loads((args.manifest_dir / f"FIG-0{i}.json").read_text(encoding="utf-8"))
        if manifest.get("publication_gate") != "PUBLICATION_READY":
            raise AssertionError(f"FIG-0{i} is not PUBLICATION_READY")
        captions[i] = manifest["caption"]

    for num in range(1, 7):
        figure_path = args.figure_dir / FIGURE_STEMS[num]
        if not figure_path.exists():
            raise FileNotFoundError(figure_path)

        cap_idx = None
        for idx, p in enumerate(doc.paragraphs):
            if p.text.strip().startswith(f"Figure {num}."):
                cap_idx = idx
                break
        if cap_idx is None:
            raise RuntimeError(f"caption Figure {num} not found")

        cap = doc.paragraphs[cap_idx]
        if cap.runs:
            cap.runs[0].text = captions[num]
            for run in cap.runs[1:]:
                run.text = ""
        else:
            cap.add_run(captions[num])

        imgp = None
        for j in range(cap_idx - 1, max(-1, cap_idx - 4), -1):
            p = doc.paragraphs[j]
            if p._p.xpath(".//w:drawing"):
                imgp = p
                break
        if imgp is None:
            raise RuntimeError(f"image paragraph for Figure {num} not found")

        for child in list(imgp._p):
            if child.tag != qn("w:pPr"):
                imgp._p.remove(child)

        run = imgp.add_run()
        run.add_picture(str(figure_path), width=Inches(args.embedded_width_in))
        imgp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
