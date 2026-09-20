from docx import Document
from copy import deepcopy
from pathlib import Path
import csv, json, shutil, difflib, re, textwrap, zipfile

ROOT=Path('/mnt/data/janus_n11_citation_novelty')
MAIN_IN=Path('/mnt/data/janus_n10_rewrite/Janus_MoSSe_ClaimFrozen_Manuscript_v1.docx')
SI_IN=Path('/mnt/data/janus_n10_rewrite/Janus_MoSSe_ClaimFrozen_SI_v1.docx')
MAIN_OUT=ROOT/'Janus_MoSSe_CitationAudited_Manuscript_v2.docx'
SI_OUT=ROOT/'Janus_MoSSe_CitationAudited_SI_v2.docx'


def set_text(p, text):
    if p.runs:
        p.runs[0].text=text
        for r in p.runs[1:]:
            r.text=''
    else:
        p.add_run(text)


def find_para(doc, starts):
    for p in doc.paragraphs:
        if p.text.startswith(starts): return p
    raise KeyError(starts)

# --- Main manuscript edits ---
doc=Document(MAIN_IN)
repls={
"Twist suppresses interfacial corrugation": "Twist suppresses interfacial corrugation in many incommensurate two-dimensional contacts, but a laterally inversion-asymmetric registry landscape can make finite-contact pinning direction dependent. We analyze the published three-shell generalized stacking-fault-energy (GSFE) landscape of 2H Janus MoSSe using an unguided finite contact whose center of mass is free in the full two-dimensional registry plane. Headline depinning thresholds are followed from the lowest-energy zero-force minimum rather than selected from an unconstrained last-surviving-minimum envelope. For the N=127 reference contact at theta=1.5 deg, the prepared rigid branch loses metastability at F_c,+*=3.071135 and F_c,-*=1.258672. Same-spectrum symmetrization removes the global directional bias and exact spatial inversion swaps the thresholds and reverses the deterministic current. Compact hexagonal and disk-like contacts collapse under theta sqrt(N), rotated triangular cuts can reverse the prepared response through registry-state switching, and zero-mean forcing produces cycle-resolved integer lattice-vector transport over all 91 sampled amplitude-period points. Representative plateaus are attracting under independent Floquet checks. Thermal exact-winding identity is fragile, whereas the full mean current vector remains statistically resolved at T*=0.65 but not at T*=0.70 under the present finite-sample protocol. Two in-plane relaxation representations preserve the directional split while renormalizing absolute thresholds. The claims are therefore DFT-anchored reduced-model results, not a parameter-free experimental device prediction.",
"Twisted van der Waals interfaces are a canonical setting": "Twisted van der Waals interfaces are a canonical setting for structural lubricity because registry forces from different parts of a finite contact partially cancel as the local phase winds across the interface [1-4]. Finite size, contact shape, boundaries, and elastic relaxation nevertheless remain active variables and can alter both the magnitude and topology of the effective sliding landscape [5-8]. Recent work continues to sharpen the role of edge geometry and moire-boundary mechanics in finite-contact friction and critical-angle scaling [25,26]. Once the local registry energy lacks inversion symmetry, the same finite-contact interference that suppresses corrugation can therefore act differently under opposite loading directions.",
"Janus transition-metal dichalcogenides provide": "Janus transition-metal dichalcogenides provide a natural materials setting for testing this possibility. Replacing the two chalcogen planes by different species produces an out-of-plane polar monolayer [9,10], while bilayer stacking and lateral sliding modify interlayer coupling and polarization [10,11,14]. Out-of-plane polarity and lateral registry-space inversion, however, are distinct symmetry statements. A defensible mechanical-rectification mechanism must begin from the actual lateral energy landscape rather than assuming that the Janus label alone implies a friction diode. The three-shell GSFE coefficient set reported for bilayer Janus MoSSe by Angeli, Schleder, and Kaxiras [12] supplies a material-anchored conservative input for such a test.",
"Directional friction and friction-diode behavior": "Directional friction, friction dissymmetry, and friction-diode behavior have already been demonstrated in two-dimensional interfaces [15,22], and rectification and phase locking on one- and two-dimensional periodic substrates are established nonlinear-transport phenomena [19,20,23,24]. First-principles-informed stochastic friction models likewise show why a static energy surface and a dynamical friction observable should not be conflated [16]. The contribution sought here is therefore narrower than a generic 'first friction diode', 'first finite-size law', or 'first ratchet' claim: we ask whether the published MoSSe GSFE supports a branch-resolved, full-2D prepared-state stability asymmetry, whether that asymmetry survives matched same-spectrum falsification and finite-contact robustness tests, and whether the same material-anchored landscape supports lattice-vector relative-periodic transport under zero-mean forcing.",
"The primary conservative input is the corrected": "The primary conservative input is the three-shell Fourier representation of the 2H Se-S-Se-S MoSSe GSFE coefficient set reported by Angeli et al. [12]. For each reciprocal shell, the six C6-related reciprocal vectors split into two C3-related triads connected by G -> -G. C3 symmetry together with the reality of the scalar registry energy makes the corresponding complex Fourier amplitudes conjugate. Pairing each G with -G and choosing one C3-related triad as the positive representatives therefore gives the equivalent real form used here,",
"with W=(7.9, 0.1, 0.1) meV": "with W=(7.9, 0.1, 0.1) meV and phi=(131.9, 1.2, 85.4) deg for the asymmetric 2H surface [12]. Choosing the conjugate reciprocal triad changes phi -> -phi, which is the spatially inverted convention; the same-spectrum inversion contracts are tested explicitly below. Energies are normalized by 6W1. In the source first-principles workflow, the aligned bilayer was sampled on a 9 x 9 lateral-displacement grid; the in-plane atomic coordinates were fixed at each registry while the interlayer distance was relaxed before the stacking-dependent quantities were fitted [12]. Ref. [13] is an erratum to Ref. [12], but its published corrections concern inverse-effective-mass labeling, Brillouin-zone path information, and a Supplemental interface label rather than the GSFE Fourier coefficients or the conjugate phase pairing used here. Thus the imported GSFE contains local registry-by-registry vertical separation relaxation, but not collective corrugation or edge reconstruction of a finite flake. The same 2022 parameterization provides three 3R configurations used below as material-derived controls.",
"Because elasticity and structural relaxation can renormalize": "Because elasticity and structural relaxation can renormalize friction even when the underlying interlayer energy surface is held fixed [7,8,18], we use two independent in-plane relaxation models without replacing the published GSFE. Tier 1 is a free-edge triangular finite-element membrane for the N=127 contact. Each site acquires a two-component internal displacement while the global center-of-mass registry and imposed twist remain external coordinates; rigid x/y translation and infinitesimal rigid rotation are projected out. The total energy is the sitewise GSFE plus isotropic membrane strain energy, and the resulting 2+(2N-3)=253 degree-of-freedom equilibrium is continued under positive- and negative-y load while monitoring the full Hessian.",
"The principal result is branch-resolved": "The principal result is branch-resolved as well as fully two-dimensional. The published MoSSe GSFE and finite-contact geometry are sufficient, within the local-GSFE reduced model, to produce a direction-split stability boundary when loading begins from the zero-force ground state. Dense/adaptive prepared-branch verification rules out a switch to a different metastable basin as the source of the headline split, while same-spectrum symmetrization and exact inversion remove or reverse the directional observables as symmetry requires. The registry-origin audit adds a complementary negative result: the physical asymmetry is invariant, but an arbitrary cosine/sine odd-sector continuation is not. Consequently the mechanism is attributed to nonremovable registry inversion asymmetry, not to a gauge-dependent eta susceptibility or optimum.",
"The theta sqrt(N) similarity is substantially": "The theta sqrt(N) similarity is substantially more robust than in the initial ten-size calculation. Larger hexagons and an independent disk-like compact family remain collapsed to about 0.1% or better across a dense Theta=0-20 grid, supporting the structure-factor interpretation for compact rigid contacts. This result sits alongside established shape- and edge-controlled finite-contact scaling [5,7] and very recent moire-boundary/edge-energy theories [25,26]; the novelty claim is therefore not the generic existence of edge-controlled scaling or an inverse-linear-size angle. Boundary orientation behaves differently in the present model: sufficiently rotated triangular cuts bring two registry minima into competition; their zero-force energy crossing changes the prepared ground state and reverses its branch-followed directional bias. The A/B competition and crossing persist across the tested surrogate outer-site weights, but the switch angle - and whether it falls inside the nominal 0-3 deg interval - is edge-model dependent. The qualitative registry competition is therefore a reduced-model prediction, whereas the numerical reversal angle is not a universal material constant.",
"A static threshold inequality does not by itself determine": "A static threshold inequality does not by itself determine a rocking-ratchet phase diagram. Generic rectification and phase locking on periodic substrates are established [19,20,23,24]; the material-specific question here is whether the same finite MoSSe contact that exhibits prepared-state directional depinning supports lattice-vector relative-periodic transport under zero-mean longitudinal forcing. The 91-point F0-tau scan is stronger than a mean-displacement classification because every sampled state repeats one winding cycle-by-cycle after transients, and representative plateaus spanning several windings are independently verified as attracting relative-periodic states. Thermal trajectories then separate four notions of robustness: exact target-winding identity is lost first; the negative-y cycle-sign bias becomes unresolved by T*=0.60; the v component of the mean current becomes unresolved at T*=0.65; and the full two-component mean vector remains statistically nonzero at 0.65 but not at 0.70. These are pointwise and joint finite-sample statements under the present protocol, not a universal thermal critical temperature or a complete operating diagram outside the tested drive, damping and reduced-temperature ranges.",
"The material conclusion can now be stated more sharply": "The material conclusion can now be stated more sharply. Translation-minimized optimization shows that the published 2H MoSSe GSFE coefficient set is genuinely inversion asymmetric in registry space, and the reciprocal-triad phase diagnostic provides an independent translation-invariant check. Matched same-2H symmetrization and exact inversion satisfy the expected static and dynamical contracts to numerical precision, which is stronger causal evidence within the reduced model than comparing only chemically different stackings. What remains unknown is the map from a physical experimental polarization change P_z to the complete set of GSFE amplitudes and phases. Exact spatial inversion is therefore a mathematical symmetry control, not a demonstrated polarization reversal. A quantitative polarity-to-rectification law requires paired polarity-resolved first-principles surfaces or equivalent atomistic input.",
"Three model-input limitations remain central": "Three model-input limitations remain central. First, the source DFT GSFE relaxes the interlayer separation independently at each aligned registry [12], but the finite-contact model still represents the interface as a sitewise local-registry sum. Collective out-of-plane corrugation or buckling, atomistic edge reconstruction and chemistry, defects, and nonlocal edge energetics are absent. The in-plane calculations show why a higher model level matters: absolute thresholds shift with compliance even when the directional ratio survives. A public empirical SW plus registry-dependent interlayer candidate is therefore retained only as a qualitative stress test because it fails the published higher-harmonic GSFE contract. Second, load-dependent registry energetics, dissipation, participating mass, and laboratory state preparation are not supplied by the present GSFE; reduced temperature and time must not be converted into kelvin or physical frequency without independent calibration. Third, physical polarity reversal is not represented by eta scaling or the matched mathematical inversion control. Beyond these material-input limitations, the isolated higher-order stability point near theta=2.9734 deg merits a dedicated two-parameter normal-form study, and the continuous amplitude-period phase diagram may contain unsampled multistability or bifurcation structure.",
"A finite contact constructed from the corrected": "A finite contact constructed from the published 2H Janus MoSSe GSFE exhibits branch-followed full two-dimensional mechanical rectification without an added transverse guide. At theta=1.5 deg the prepared N=127 rigid thresholds are F_c,+*=3.071135 and F_c,-*=1.258672, and dense/adaptive continuation over 0-3 deg finds no exchange between the prepared ground family and the competing metastable family. Translation-minimized asymmetry diagnostics and matched same-2H controls identify registry-space inversion asymmetry as the causal conservative ingredient within the reduced model. Compact-contact thresholds obey theta sqrt(N) similarity with finite-size corrections toward the inverse-linear-size law, while rotated triangular boundaries can reverse the prepared bias through a ground-state registry switch whose angle is edge-model dependent. Zero-mean forcing generates cycle-resolved integer lattice-vector transport over all 91 sampled amplitude-period points, with representative plateaus independently verified as attracting. Thermal exact-winding identity is fragile, but the full mean vector current remains resolved at T*=0.65 and becomes unresolved at T*=0.70 under the present finite-sample protocol. Linear FEM and independently reconstructed nonlinear VFF relaxations soften absolute thresholds while preserving the directional split. The resulting evidence supports a self-consistent DFT-anchored reduced-model mechanism; it does not establish a unique experimental friction coefficient, a physical polarity law, full 3D atomistic relaxation, or a universal nonlinear phase diagram.",
"The accompanying numerical materials contain": "The accompanying numerical materials contain the canonical unguided rigid-model release, the linear elastic-validation package, the independently reconstructed nonlinear bond-angle source and outputs, the Tier-2 interlayer-potential prescreen, and the Gate-level audit records used to freeze manuscript claims. The rigid and audit packages include the published-GSFE engine; dense prepared-branch enumeration and continuation; matched symmetry and exact-inversion controls; registry-origin and odd-sector gauge audits; exact finite-size/shape and boundary analyses; cycle-resolved vector-locking and Floquet checks; trajectory-level thermal data and joint-vector bootstrap summaries; and elasticity validation outputs. A citation-audit patch removes an earlier source-comment attribution that incorrectly described the 2024 Erratum as a GSFE phase-convention correction; no numerical coefficient or computed result changes under that documentation correction. No archival DOI is claimed until external repository deposition is completed."
}

# Paragraph 3 (abstract) matches generic start too; target by index explicitly first.
set_text(doc.paragraphs[3], repls["Twist suppresses interfacial corrugation"])
for key,new in repls.items():
    if key=="Twist suppresses interfacial corrugation": continue
    p=find_para(doc,key); set_text(p,new)

# Replace any residual provenance wording in remaining body paragraphs and tables.
for p in doc.paragraphs:
    if 'corrected published' in p.text:
        set_text(p,p.text.replace('corrected published','published'))
    if 'corrected higher-harmonic' in p.text:
        set_text(p,p.text.replace('corrected higher-harmonic','published higher-harmonic'))
    if 'corrected DFT reference' in p.text:
        set_text(p,p.text.replace('corrected DFT reference','published DFT reference'))
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                txt=p.text
                txt2=txt.replace('corrected DFT','published DFT').replace('corrected published','published').replace('corrected GSFE','published GSFE')
                if txt2!=txt: set_text(p,txt2)

# Add new references at end, preserving existing numbering.
refs=[
"22. Lu, L.; Ying, T.; Cui, C.-X.; Yang, C.; Leng, J.; Li, J.; Jiang, J.-W.; Chang, T. Friction Dissymmetry on Hexagonal Boron Carbon Nitride. Nano Letters 25, 7909-7915 (2025). https://doi.org/10.1021/acs.nanolett.5c01360.",
"23. Reichhardt, C.; Olson, C. J.; Hastings, M. B. Rectification and Phase Locking for Particles on Symmetric Two-Dimensional Periodic Substrates. Physical Review Letters 89, 024101 (2002). https://doi.org/10.1103/PhysRevLett.89.024101.",
"24. Wang, C.-L.; Tekic, J.; Duan, W.-S.; Shao, Z.-G.; Yang, L. Ratchet effect and amplitude dependence of phase locking in a two-dimensional Frenkel-Kontorova model. The Journal of Chemical Physics 138, 034307 (2013). https://doi.org/10.1063/1.4776226.",
"25. Liu, Z. Edge energy fluctuation: A unified theory of friction scaling laws in twisted layered material interfaces. arXiv:2609.08063v1 (2026). Preprint submitted 8 September 2026. https://arxiv.org/abs/2609.08063.",
"26. Wang, K.; Chen, A.; Huang, J.; Zhang, Y.; Liang, Y.; Han, Q. Moire boundary dominated twisting graphene friction: Scaling laws and geometrical control. International Journal of Solids and Structures 337, 114061 (2026). https://doi.org/10.1016/j.ijsolstr.2026.114061."
]
for ref in refs:
    p=doc.add_paragraph(ref)
    p.style='Normal'

doc.save(MAIN_OUT)

# --- SI edits ---
si=Document(SI_IN)
for p in si.paragraphs:
    txt=p.text
    txt2=(txt.replace('corrected three-shell Fourier GSFE','published three-shell Fourier GSFE')
              .replace('same corrected 2H GSFE','same published 2H GSFE')
              .replace('corrected 2H GSFE','published 2H GSFE')
              .replace('corrected DFT reference','published DFT reference')
              .replace('corrected higher-harmonic','published higher-harmonic')
              .replace('corrected GSFE','published GSFE'))
    if txt2!=txt: set_text(p,txt2)
for table in si.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                txt=p.text
                txt2=txt.replace('corrected DFT','published DFT').replace('corrected 2H GSFE','published 2H GSFE').replace('corrected GSFE','published GSFE')
                if txt2!=txt: set_text(p,txt2)
# Insert explicit provenance note into scope paragraph.
p=find_para(si,'The conservative material input is the published')
set_text(p, "The conservative material input is the published three-shell Fourier GSFE coefficient set for 2H MoSSe used in the main text. The real three-vector-per-shell implementation is obtained by pairing the two C3-related reciprocal triads required to be complex conjugates by the reality of the scalar energy; choosing the conjugate triad corresponds to spatial inversion. The 2024 Angeli Erratum does not modify these GSFE coefficients or this phase pairing; it corrects effective-mass labeling, band-path information, and one Supplemental interface label. A rigid finite contact is formed by summing the local GSFE over lattice sites after twist. Unless otherwise stated, the reference contact is the N=127 compact hexagon, and all static headline results use the lowest-energy zero-force minimum as the prepared state. The equations of motion use m*=1 and gamma*=4. Angles are in degrees and forces, energies, temperature, and time are reduced quantities unless an explicit unit is shown.")
si.save(SI_OUT)

# --- Code provenance correction copies and patch ---
srcdir=Path('/mnt/data/janus_baseline_gate/Janus_MoSSe_CanonicalUnguided_BaselineRelease/source')
codeout=ROOT/'source_provenance_corrected'; codeout.mkdir(exist_ok=True)
patch_sections=[]
for name in ['janus_fourier_landscapes_v12.py','dft_guided_material_v12.py']:
    old=(srcdir/name).read_text()
    new=old
    if name=='janus_fourier_landscapes_v12.py':
        old_head='''The literature parameterization follows Angeli, Schleder & Kaxiras,\nPhys. Rev. B 106, 235159 (2022), with the phase convention corrected in\nPhys. Rev. B 109, 199902(E) (2024).  The GSFE is represented as\n\n  Omega(r) = sum_l sum_j W_l exp[i g_j^l.r + i(-1)^j phi_l].\n\nPairing +/- reciprocal vectors and dividing by E0=6 W1 gives a real,\ndimensionless three-vector representation per shell:\n'''
        new_head='''The literature parameterization follows the three-shell GSFE coefficient set\nreported by Angeli, Schleder & Kaxiras, Phys. Rev. B 106, 235159 (2022).\nFor each shell, C3 symmetry and the reality of the scalar registry energy make\nthe two C3-related reciprocal-vector triads complex conjugates. Choosing one\ntriad as the positive representatives and pairing each G with -G gives the real\nform below. Choosing the conjugate triad corresponds to spatial inversion.\nThe 2024 Erratum (Phys. Rev. B 109, 199902) does not change these GSFE\ncoefficients or this conjugate phase pairing; its corrections concern effective-\nmass labeling, Brillouin-zone path information, and one Supplemental interface\nlabel. Dividing by E0=6 W1 gives a real, dimensionless three-vector\nrepresentation per shell:\n'''
        new=new.replace(old_head,new_head)
        new=new.replace('# Published GSFE coefficients from Table II of Angeli et al. (2022),\n# using the 2024 erratum phase convention. Only configurations used in v12\n', '# Published GSFE coefficients from Table II of Angeli et al. (2022).\n# The 2024 Erratum does not modify these GSFE coefficients or phase pairing.\n# Only configurations used in v12\n')
    else:
        new=new.replace('Uses published MoSSe GSFE Fourier amplitudes/phases from Angeli et al.\n(PRB 106, 235159, 2022) with the 2024 erratum convention.  The DFT data\n', 'Uses the published MoSSe GSFE Fourier amplitudes/phases from Angeli et al.\n(PRB 106, 235159, 2022). The real implementation pairs conjugate reciprocal\ntriads as required by C3 symmetry and a real scalar energy. The 2024 Erratum\ndoes not alter the GSFE coefficients or phase pairing. The DFT data\n')
    (codeout/name).write_text(new)
    diff=''.join(difflib.unified_diff(old.splitlines(True),new.splitlines(True),fromfile=str(srcdir/name),tofile=str(codeout/name)))
    patch_sections.append(diff)
(ROOT/'angeli_phase_provenance_correction.patch').write_text('\n'.join(patch_sections))

# --- Citation audit CSV ---
rows=[
(1,'Wang et al. 2024 RMP','2D sliding/pinning; boundaries/elasticity','VERIFIED','Keep','Review supports broad framework; not novelty source.'),
(2,'Koren & Duerig 2016 PRB','finite-contact moire cancellation/scaling','VERIFIED','Keep','Supports finite-contact interference/scaling.'),
(3,'Yan et al. 2023 JMPS','moire-boundary origin of finite-contact scaling','VERIFIED','Keep','Supports boundary/scaling context.'),
(4,'Verhoeven et al. 2004 PRB','rigid finite-flake size/shape/pulling direction','VERIFIED','Keep','Direct model precedent.'),
(5,'Yan et al. 2024 JMPS','shape/edge-dependent scaling','VERIFIED','Keep','Direct overlap; narrows novelty.'),
(6,'Liao et al. 2022 Nat Mater','edge pinning in vdW contacts','VERIFIED','Keep','Supports boundary importance.'),
(7,'Gao et al. 2025 ACS Nano','elastic edge/corner moire pinning','VERIFIED','Keep','Supports elasticity and boundary sensitivity.'),
(8,'Minkin et al. 2025 PR Materials','relaxation changes PES corrugation/topology','VERIFIED','Keep','Strong limitation/support for elastic robustness motivation.'),
(9,'Lu et al. 2017 Nat Nanotech','Janus TMD material basis','VERIFIED','Keep','Supports Janus monolayer basis.'),
(10,'Zhang et al. 2020 JACS','polar Janus MoSSe interlayer coupling','VERIFIED','Keep','Supports polarity/interlayer coupling.'),
(11,'Lin et al. 2024 Nanoscale','sliding ferroelectricity/moire in bilayer MoSSe','VERIFIED','Keep','Supports sliding/stacking/polarization context; not tribological depinning.'),
(12,'Angeli et al. 2022 PRB','GSFE coefficients and DFT workflow','VERIFIED','Keep - primary','Primary source for coefficients/workflow. C3+reality justify conjugate triad pairing.'),
(13,'Angeli et al. 2024 Erratum','erratum scope','VERIFIED WITH RESTRICTION','Keep only for explicit erratum note','Does NOT correct GSFE phase convention. Earlier attribution was wrong and is removed.'),
(14,'Li et al. 2022 Friction','Janus TMD sliding/shear under field/load','VERIFIED','Keep','Supports Janus tribology context.'),
(15,'Chen et al. 2026 ACS AMI','intrinsic friction diodes from asymmetric PES','VERIFIED','Keep','Closest friction-diode prior art; generic diode novelty not available.'),
(16,'Torche et al. 2022 Adv Mater Interfaces','first-principles PES + stochastic friction dynamics','VERIFIED','Keep','Supports distinction between static PES and dynamic friction.'),
(17,'Guo 2018 PCCP','MoSSe elastic constants','VERIFIED','Keep','Supports C11/C12 input.'),
(18,'Dong et al. 2023 JPCL','stiffness alters friction at fixed ISES','VERIFIED','Keep','Supports elasticity sensitivity.'),
(19,'Hanggi & Marchesoni 2009 RMP','rocked ratchet framework','VERIFIED','Keep','Generic ratchet prior art.'),
(20,'Barbi & Salerno 2000 PRE','phase locking/current reversal','VERIFIED','Keep','Generic deterministic ratchet prior art.'),
(21,'Shaidu et al. 2025 npj Comput Mater','dispersion-aware MLIP for multilayer TMDs','VERIFIED SCOPE','Keep','Prospective route only; not Janus-MoSSe validation.'),
(22,'Lu et al. 2025 Nano Letters','friction dissymmetry on graphene/h-BCN','NEW - VERIFIED','Add','Direct directional-friction prior art; further narrows generic diode novelty.'),
(23,'Reichhardt et al. 2002 PRL','2D periodic-substrate rectification/phase locking','NEW - VERIFIED','Add','Direct generic 2D mode-locking prior art.'),
(24,'Wang et al. 2013 JCP','2D Frenkel-Kontorova ratchet/phase locking','NEW - VERIFIED','Add','Additional generic 2D ratchet/locking precedent.'),
(25,'Ze Liu 2026 arXiv:2609.08063','edge-energy fluctuation/friction scaling','NEW - PREPRINT VERIFIED','Add as preprint','Very recent overlap with edge orientation/scaling; not prepared-state Janus depinning.'),
(26,'Wang et al. 2026 IJSS','moire-boundary friction/critical-angle scaling','NEW - VERIFIED','Add','Recent critical-angle/geometric scaling overlap; rotational focus.'),
]
with (ROOT/'CITATION_SUPPORT_AUDIT.csv').open('w',newline='') as f:
    w=csv.writer(f); w.writerow(['ref','source','claim_role','status','action','notes']); w.writerows(rows)

# --- Novelty matrix ---
nov=[
('Generic friction diode / opposite-direction friction','NOT NOVEL','Chen 2026; Lu 2025','Do not use first friction diode / first directional friction.'),
('Asymmetric PES as source of directional friction','NOT NOVEL','Chen 2026; Lu 2025','Frame as material-specific symmetry mechanism, not generic discovery.'),
('Finite-contact twist/shape/edge scaling','NOT NOVEL','Koren 2016; Yan 2023/2024; Gao 2025; Liu 2026 preprint; Wang 2026','Use as context; specific prepared-threshold collapse is a result, not generic novelty.'),
('Inverse-linear-size critical-angle scaling','NOT NOVEL / ADJACENT PRIOR','Wang 2026; finite-size moire literature','Do not sell alpha~0.5036 or N^-1/2 as new exponent/law.'),
('Generic 2D ratchet / phase locking','NOT NOVEL','Reichhardt 2002; Wang et al. 2013; Hanggi 2009; Barbi 2000','Do not use first ratchet / first mode locking.'),
('Published MoSSe GSFE + unguided prepared full-2D directional stability boundary','LIKELY NOVEL','No direct match in >300 result slots across >30 targeted searches','Use cautious contribution wording; no first claim.'),
('Same-spectrum symmetrization + exact-inversion falsification tied to the MoSSe GSFE','LIKELY NOVEL','No direct tribology match found','Strong methodological contribution within reduced model.'),
('Ground-state registry-switch reversal for rotated finite Janus boundary','UNCERTAIN to LIKELY','Edge orientation and stacking instabilities exist in prior art; no same mechanism found','Describe as specific reduced-model mechanism, not first edge reversal.'),
('Zero-mean longitudinal rocking -> lattice-vector relative-periodic/Floquet states in same finite MoSSe contact','LIKELY NOVEL','Generic 2D locking exists, but no direct vdW/Janus finite-interface counterpart found','Emphasize material-specific integration with same static landscape.'),
('Thermal vector-current bootstrap hierarchy','METHOD/ROBUSTNESS, not standalone novelty','Statistical validation of this model','Use as robustness evidence.'),
('FEM/VFF in-plane compliance preservation of directional split','METHOD/ROBUSTNESS, not standalone novelty','Elasticity-friction prior art extensive','Use as falsification/robustness evidence.'),
('Overall integrated mechanism package','LIKELY NOVEL','No single source found combining all components','Safest novelty statement; explicitly time-stamped to 2026-09-19 search.'),
]
with (ROOT/'NOVELTY_REVERSE_SEARCH_MATRIX.csv').open('w',newline='') as f:
    w=csv.writer(f); w.writerow(['component','novelty_status','closest_prior_art','allowed_positioning']); w.writerows(nov)

# --- Bibliography additions ---
(ROOT/'NEW_REFERENCES_22_26.txt').write_text('\n'.join(refs)+'\n')
(ROOT/'NEW_REFERENCES_22_26.bib').write_text(r'''@article{Lu2025FrictionDissymmetry,
  author = {Lu, Leiling and Ying, Tianquan and Cui, Chuan-Xin and Yang, Chen and Leng, Jiantao and Li, Jianxin and Jiang, Jin-Wu and Chang, Tienchong},
  title = {Friction Dissymmetry on Hexagonal Boron Carbon Nitride},
  journal = {Nano Letters},
  volume = {25},
  pages = {7909--7915},
  year = {2025},
  doi = {10.1021/acs.nanolett.5c01360}
}
@article{Reichhardt2002Rectification,
  author = {Reichhardt, C. and Olson, C. J. and Hastings, M. B.},
  title = {Rectification and Phase Locking for Particles on Symmetric Two-Dimensional Periodic Substrates},
  journal = {Physical Review Letters},
  volume = {89},
  pages = {024101},
  year = {2002},
  doi = {10.1103/PhysRevLett.89.024101}
}
@article{Wang2013Ratchet,
  author = {Wang, Cang-Long and Tekic, Jasmina and Duan, Wen-Shan and Shao, Zhi-Gang and Yang, Lei},
  title = {Ratchet effect and amplitude dependence of phase locking in a two-dimensional Frenkel-Kontorova model},
  journal = {The Journal of Chemical Physics},
  volume = {138},
  pages = {034307},
  year = {2013},
  doi = {10.1063/1.4776226}
}
@article{Liu2026EdgeEnergyFluctuation,
  author = {Liu, Ze},
  title = {Edge energy fluctuation: A unified theory of friction scaling laws in twisted layered material interfaces},
  journal = {arXiv preprint arXiv:2609.08063},
  year = {2026},
  note = {v1, submitted 8 September 2026}
}
@article{Wang2026MoireBoundary,
  author = {Wang, Kejing and Chen, Ao and Huang, Jianzhang and Zhang, Yajiu and Liang, Yingjing and Han, Qiang},
  title = {Moire boundary dominated twisting graphene friction: Scaling laws and geometrical control},
  journal = {International Journal of Solids and Structures},
  volume = {337},
  pages = {114061},
  year = {2026},
  doi = {10.1016/j.ijsolstr.2026.114061}
}
''')

# --- Provenance correction report ---
prov='''# Angeli GSFE provenance correction\n\n## Decision\n\n**CORRECTION REQUIRED AND APPLIED.** An earlier source comment and manuscript shorthand attributed the Fourier phase convention to the 2024 Angeli Erratum. The published Erratum does not make that correction. It corrects inverse-effective-mass table labeling/caption information, Brillouin-zone path information, and one Supplemental interface label.\n\n## What remains valid\n\nThe numerical GSFE implementation is unchanged. The 2022 paper states that the scalar continuum potentials are represented on six reciprocal vectors per shell and that C3 symmetry plus reality reduce the shell coefficients to an amplitude and phase. The six vectors split into two C3 triads related by G -> -G. For a real scalar energy, the complex coefficients of opposite reciprocal vectors are conjugates. Choosing one triad as the positive representatives therefore yields\n\n    Omega_l(r) = 2 W_l sum_{m=1}^3 cos(G_lm . r + phi_l),\n\nand after normalization by E0 = 6 W1,\n\n    u_l(r) = [W_l/(3 W1)] sum_{m=1}^3 cos(G_lm . r + phi_l).\n\nThis is exactly the real three-vector-per-shell form implemented by the production code. Choosing the conjugate triad corresponds to phi_l -> -phi_l, i.e. spatial inversion. The exact-inversion gate already verifies that this operation swaps directional thresholds/current as symmetry requires.\n\n## Source roles after correction\n\n- **Angeli et al. 2022, PRB 106, 235159:** primary source for the GSFE Fourier form, coefficients, DFT registry sampling, and relaxation protocol.\n- **Angeli et al. 2024, PRB 109, 199902(E):** cited only to state the actual erratum scope; it is not the source of the phase pairing used in code.\n\n## Numerical consequence\n\nNone. No W_l, phi_l, reciprocal vector, normalization, threshold, trajectory, or figure datum is changed. This is a provenance/documentation correction, not a numerical model correction.\n'''
(ROOT/'ANGELI_GSFE_PROVENANCE_CORRECTION.md').write_text(prov)

# --- Final audit report ---
report='''# N11 Final Citation + Novelty Reverse Audit\n\n**Audit date:** 2026-09-19\n\n## Gate decision\n\n**PASS AFTER CORRECTION.** Citation metadata and source-claim alignment are adequate after correcting the Angeli-Erratum provenance statement and adding the closest recent directional-friction, 2D mode-locking, and 2026 edge/scaling prior art. The overall integrated contribution remains **LIKELY NOVEL**, not externally proven unique.\n\n## 1. Critical provenance correction\n\nThe 2024 Angeli Erratum does **not** correct the GSFE Fourier phase convention. The v1 manuscript/source shorthand implying that it did has been removed. The real paired Fourier form used numerically is instead justified from the 2022 paper's six-vector representation, C3 symmetry, and the reality condition, which require conjugate coefficients on G and -G. Choosing the conjugate C3 triad is spatial inversion. No numerical result changes.\n\n## 2. Citation audit outcome\n\nThe existing 21-reference list contains no identified fabricated paper or DOI. Ref. 13 is valid but its role is narrowed to the actual Erratum. Ref. 21 remains prospective MLIP context and is not presented as Janus-MoSSe atomistic validation. Five references are added to close important prior-art gaps:\n\n- 2025 Nano Letters: opposite-direction friction dissymmetry from a dissymmetric 2D-material PES.\n- 2002 PRL and 2013 JCP: rectification/phase locking on 2D periodic substrates.\n- 2026 arXiv preprint: edge-energy-fluctuation theory for twisted layered friction scaling.\n- 2026 IJSS: moire-boundary geometric/critical-angle scaling in graphene.\n\n## 3. Novelty reverse search\n\nThe search covered **more than 300 result slots across more than 30 targeted workstreams**, emphasizing 2024-01-01 through 2026-09-19 and adding older foundational literature. Search channels included exact-title/DOI verification, Janus MoSSe sliding/twist, friction-diode/dissymmetry work, finite-contact scaling/edge mechanics, prepared/depinning terminology, and 2D ratchet/mode-locking work. Literature search cannot prove nonexistence, so absence claims are not made.\n\n### Findings\n\n- Generic friction-diode/directional-friction behavior is **NOT NOVEL**.\n- Asymmetric PES as the origin of directional friction is **NOT NOVEL**.\n- Generic finite-contact shape/edge/twist scaling and inverse-linear-size angular scaling are **NOT NOVEL**.\n- Generic 2D ratchet/phase-locking phenomena are **NOT NOVEL**.\n- No direct match was found for the full combination of the published MoSSe GSFE, unguided prepared-ground-state full-2D stability boundaries, same-spectrum symmetry falsification, finite-contact geometry tests, and zero-mean lattice-vector relative-periodic/Floquet transport. The integrated contribution is therefore **LIKELY NOVEL** as of the search date, with no 'first' language used.\n\n## 4. Closest recent threats to novelty\n\n1. Chen et al. (2026), ACS Applied Materials & Interfaces: intrinsic friction diodes from asymmetric PES and dynamic pathway selection.\n2. Lu et al. (2025), Nano Letters: friction dissymmetry for graphene/h-BCN from a dissymmetric PES.\n3. Yan et al. (2024), JMPS and Gao et al. (2025), ACS Nano: shape/edge/twist and elastic edge/corner control of finite-contact friction.\n4. Liu (2026), arXiv:2609.08063: edge-energy fluctuation as a unified friction-scaling mechanism.\n5. Wang et al. (2026), IJSS 337, 114061: moire-boundary scaling with critical angle proportional to inverse size in a rotational graphene setting.\n6. Reichhardt et al. (2002) and Wang et al. (2013): generic 2D periodic-substrate ratchet/phase-locking precedents.\n\n## 5. Manuscript language frozen after audit\n\nAllowed: material-specific, reduced-model, branch-resolved, same-spectrum-falsified, tested-scope claims.\n\nDisallowed: 'first friction diode', 'first mode locking', 'new universal N^-1/2 exponent', physical eta susceptibility/optimum, exact edge-switch angle as a material constant, or Erratum-corrected GSFE phase convention.\n\n## 6. Next gate\n\n**Red Team.** The next step should assume the paper's central conclusion is wrong and attack hidden assumptions, confounders, statistical artifacts, model incompleteness, novelty overlap, and reproducibility before another five-reviewer pass.\n'''
(ROOT/'N11_FINAL_CITATION_NOVELTY_AUDIT_REPORT.md').write_text(report)

# --- Package ---
zip_path=ROOT/'N11_CitationNovelty_AuditBundle.zip'
files=[MAIN_OUT,SI_OUT,ROOT/'CITATION_SUPPORT_AUDIT.csv',ROOT/'NOVELTY_REVERSE_SEARCH_MATRIX.csv',ROOT/'ANGELI_GSFE_PROVENANCE_CORRECTION.md',ROOT/'NEW_REFERENCES_22_26.txt',ROOT/'NEW_REFERENCES_22_26.bib',ROOT/'angeli_phase_provenance_correction.patch',ROOT/'N11_FINAL_CITATION_NOVELTY_AUDIT_REPORT.md',codeout/'janus_fourier_landscapes_v12.py',codeout/'dft_guided_material_v12.py']
with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED) as z:
    for f in files:
        z.write(f, f.relative_to(ROOT) if f.is_relative_to(ROOT) else f.name)
print('wrote',MAIN_OUT,SI_OUT,zip_path)
