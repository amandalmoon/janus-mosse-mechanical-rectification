from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src='/mnt/data/janus_n17_upstream_strengthening/Janus_MoSSe_UpstreamStrengthened_SI_v5.docx'
out='/mnt/data/janus_n17_upstream_strengthening/Janus_MoSSe_UpstreamStrengthened_SI_v5_fixed.docx'
doc=Document(src)

def prevent_row_split(table):
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn('w:cantSplit')) is None:
            trPr.append(OxmlElement('w:cantSplit'))
    trPr = table.rows[0]._tr.get_or_add_trPr()
    if trPr.find(qn('w:tblHeader')) is None:
        trPr.append(OxmlElement('w:tblHeader'))

# Table S8a starts on a fresh page; caption stays with table.
p_s8a = next(p for p in doc.paragraphs if p.text.startswith('Table S8a.'))
p_s8a.paragraph_format.page_break_before = True
p_s8a.paragraph_format.keep_with_next = True
p_s8b = next(p for p in doc.paragraphs if p.text.startswith('Table S8b.'))
p_s8b.paragraph_format.keep_with_next = True

# Move S10 caption before its table.
p_s10cap = next(p for p in doc.paragraphs if p.text.startswith('Table S10.'))
t_s10 = next(t for t in doc.tables if t.cell(0,0).text.strip()=='case' and any(c.text.strip()=='vs FEM F+' for c in t.rows[0].cells))
p_s10cap._p.getparent().remove(p_s10cap._p)
t_s10._tbl.addprevious(p_s10cap._p)
p_s10cap.paragraph_format.keep_with_next = True

# Keep S10 heading + explanatory paragraph together if possible.
plist = list(doc.paragraphs)
for i,p in enumerate(plist):
    if p.text.startswith('S10. Nonlinear'):
        p.paragraph_format.keep_with_next = True
        if i+1 < len(plist):
            plist[i+1].paragraph_format.keep_with_next = True
        break

for t in doc.tables:
    prevent_row_split(t)

doc.save(out)
print(out)
