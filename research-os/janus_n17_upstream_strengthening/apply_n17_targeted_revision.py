from pathlib import Path
from docx import Document
import shutil, json, pandas as pd
from openpyxl import load_workbook

BASE=Path('/mnt/data/janus_n15_targeted_revision')
OUT=Path('/mnt/data/janus_n17_upstream_strengthening')
main_in=BASE/'Janus_MoSSe_RedTeamScoped_Manuscript_v4.docx'
si_in=BASE/'Janus_MoSSe_RedTeamScoped_SI_v4.docx'
main_out=OUT/'Janus_MoSSe_UpstreamStrengthened_Manuscript_v5.docx'
si_out=OUT/'Janus_MoSSe_UpstreamStrengthened_SI_v5.docx'
shutil.copy2(main_in,main_out);shutil.copy2(si_in,si_out)

def replace_para(doc, starts, newtext):
    for p in doc.paragraphs:
        if p.text.startswith(starts):
            p.text=newtext
            return True
    return False

def append_to_para(doc, starts, extra):
    for p in doc.paragraphs:
        if p.text.startswith(starts):
            p.add_run(extra)
            return True
    return False

# MAIN
D=Document(main_out)
# Methods 2.1 source-input scope
append_to_para(D,
 'with W=(7.9, 0.1, 0.1) meV',
 ' Because the source table reports W_l and phi_l only to 0.1 meV and 0.1 deg and does not provide a coefficient covariance or machine-readable 9 x 9 total-energy grid in the accessed primary-source materials, the published coefficients are treated as the model definition rather than as statistically calibrated random variables. A separate input-resolution audit perturbs every reported W_l and phi_l across its half-unit rounding interval; that test bounds sensitivity to publication precision but is not a substitute for source DFT/fit uncertainty.'
)
# Discussion 4.3 joint damping-drive interaction
old='A static threshold inequality does not by itself determine a rocking-ratchet phase diagram.'
for p in D.paragraphs:
    if p.text.startswith(old):
        txt=p.text
        anchor='The sampled vector-locking map is therefore a property of the declared m*=1, gamma*=4 reduced dynamical protocol, not a material-only MoSSe phase diagram.'
        extra=(' A compact 72-point interaction audit over gamma*=2, 3, 4, 5, 6 and 8, F0*=1.4, 2.0, 2.6 and 3.0, and tau*=30, 40 and 60 reinforces that boundary: only 1 of the 12 tested drive points preserves the gamma*=4 winding across every tested damping value, while a single drive point can exhibit as many as five distinct winding pairs across damping.')
        if extra.strip() not in txt:
            txt=txt.replace(anchor,anchor+extra)
        p.text=txt
        break
# Discussion 4.5 input precision + unavailable fit covariance
for p in D.paragraphs:
    if p.text.startswith('Three model-input limitations remain central.'):
        txt=p.text
        anchor='Collective out-of-plane corrugation or buckling, atomistic edge reconstruction and chemistry, defects, and nonlocal edge energetics are absent.'
        extra=(' The published-precision audit also shows that varying every tabulated W_l and phi_l within the half-unit interval implied by their reported decimal precision preserves the positive static split and the canonical (1,-1) winding; across 16 joint Latin-hypercube checks, rho remains between 0.41716 and 0.42002. This bounds sensitivity to coefficient reporting precision only. The original fit covariance and raw 9 x 9 total-energy table are not available in the accessed source materials, so DFT/fit uncertainty is not propagated and the quantitative predictions remain conditional on the published Fourier parameterization.')
        if extra.strip() not in txt:
            txt=txt.replace(anchor,anchor+extra)
        p.text=txt
        break
# Thermal Results precision justification
for p in D.paragraphs:
    if p.text.startswith('The stationarity audit establishes the long-window thermal inference protocol.'):
        txt=p.text
        anchor='At T*=0.70 a 50,000-resample trajectory bootstrap gives u in [0.01265,0.06132] and v in [-0.10088,-0.05195], and the zero-current vector remains outside the 95% joint region.'
        extra=(' A 500-resample without-replacement trajectory-count audit provides an explicit precision check: the fraction of subsamples whose 95% Hotelling test excludes the zero vector reaches at least 0.95 by N=100, 200, 200, 300 and 500 at T*=0.50, 0.55, 0.60, 0.65 and 0.70, respectively; the pooled N=1000 design is therefore conservative relative to the weakest sampled signal.')
        if extra.strip() not in txt:
            txt=txt.replace(anchor,anchor+extra)
        p.text=txt
        break
# Data/code availability list
for p in D.paragraphs:
    if p.text.startswith('The accompanying numerical materials contain'):
        p.add_run(' The audit release also includes the publication-resolution GSFE coefficient sensitivity, the joint damping-drive interaction map, and the trajectory-count precision/convergence analysis added after independent review.')
        break
D.save(main_out)

# SI
S=Document(si_out)
for p in S.paragraphs:
    if p.text.startswith('The conservative material input is the published three-shell Fourier GSFE'):
        p.add_run(' The source table reports W_l and phi_l to one decimal place but does not supply a fit covariance or machine-readable 9 x 9 total-energy table in the accessed primary-source materials. We therefore treat the published Fourier coefficients as the model definition. A publication-resolution audit varies all six tabulated amplitudes/phases within +/-0.05 meV or +/-0.05 deg, respectively: a first-order hyperrectangle bound gives Delta F_c*=1.75899-1.86594 and rho=0.41641-0.42079, while 16 joint Latin-hypercube evaluations give Delta F_c*=1.77264-1.84708 and rho=0.41716-0.42002. All 16 retain positive rectification and the canonical deterministic checks retain the (1,-1) winding. This is a reporting-precision bound, not an estimate of the unavailable source DFT/fit covariance.')
        break
for p in S.paragraphs:
    if p.text.startswith('The canonical (1,-1) leading multiplier agrees'):
        p.add_run(' A separate 72-point joint damping-drive audit spans gamma*=2, 3, 4, 5, 6 and 8; F0*=1.4, 2.0, 2.6 and 3.0; and tau*=30, 40 and 60. Every sampled point is cycle locked, but only 1 of 12 drive points preserves the gamma*=4 winding across all six damping values. The median fraction of damping values retaining the gamma*=4 winding is 0.417, and up to five distinct winding pairs occur at one fixed drive point. Thus damping does not merely shift a plateau boundary; it reorganizes the sampled winding map.')
        break
for p in S.paragraphs:
    if p.text.startswith('Each trajectory contains 60 burn cycles and 100 measured cycles at dt=0.02.'):
        p.add_run(' To quantify trajectory-count precision, 500 without-replacement subsamples are drawn from each pooled N=1000 high-temperature ensemble at N=50, 100, 200, 300, 500 and 750, with the full N=1000 result retained as the target. The minimum N at which at least 95% of subsamples exclude the zero vector by the same 95% Hotelling criterion is 100 at T*=0.50, 200 at 0.55, 200 at 0.60, 300 at 0.65 and 500 at 0.70. This convergence audit motivates the two-seed N=1000 high-temperature analysis as a precision choice rather than a post hoc significance boundary.')
        break
for p in S.paragraphs:
    if p.text.startswith('The audit resolves the computational ambiguities'):
        p.add_run(' A further source-input limitation is that the original GSFE fit covariance/raw 9 x 9 energy table is not available in the accessed primary-source materials; only sensitivity to the published decimal reporting precision can be bounded here.')
        break
for p in S.paragraphs:
    if p.text.startswith('The accompanying materials retain the canonical unguided rigid-model release'):
        p.add_run(' The post-review package additionally contains the publication-resolution GSFE audit, the 72-point joint damping-drive map, and the thermal trajectory-count convergence analysis.')
        break
S.save(si_out)

# Update Claim-Evidence Matrix v3 (CSV/MD/XLSX)
mat_in=Path('/mnt/data/janus_n13_claim_figure_revision/CLAIM_EVIDENCE_MATRIX_v2.csv')
df=pd.read_csv(mat_in)

def upd(cid, **kw):
    idx=df.index[df.ID==cid]
    if len(idx)!=1: raise RuntimeError(cid)
    i=idx[0]
    for k,v in kw.items(): df.at[i,k]=v

m0=df[df.ID=='M0'].iloc[0]
upd('M0',
    **{'Quantitative / Falsification Contract':str(m0['Quantitative / Falsification Contract'])+' Publication-resolution propagation: W_l +/-0.05 meV and phi_l +/-0.05 deg; 16 joint LHS checks give DeltaFc*=1.77264-1.84708 and rho=0.41716-0.42002; canonical winding remains (1,-1). Actual source fit covariance/raw-grid uncertainty is not available.',
       'Scope':'Published Fourier parameterization. The added perturbation audit bounds decimal reporting precision only; actual DFT/fit uncertainty is NOT_ASSESSABLE from the accessed source data.',
       'Recommended Manuscript Wording':'Use the 2022 coefficients as the declared model input; state that publication-resolution perturbations preserve the result but do not substitute for an unavailable source fit covariance/raw GSFE grid.',
       'Forbidden / Retired Wording':str(m0['Forbidden / Retired Wording'])+' Do not call the reporting-resolution ensemble a DFT uncertainty distribution.',
       'Evidence Artifact':'janus_n17_upstream_strengthening/fr_m01_published_precision_summary.json',
       'Manuscript Action':'UPDATED after Five Reviewer FR-M01: source-limited uncertainty handled by explicit scope + reporting-precision bound.'})

c13=df[df.ID=='C13'].iloc[0]
upd('C13',
    **{'Quantitative / Falsification Contract':str(c13['Quantitative / Falsification Contract'])+' Joint damping-drive audit: 72 points over gamma*=2,3,4,5,6,8; F0*=1.4,2.0,2.6,3.0; tau*=30,40,60. Only 1/12 drive points preserves the gamma*=4 winding for all tested gamma; median preservation fraction 0.417; up to 5 winding pairs at one drive point.',
       'Scope':'The 91-point headline map is strictly for m*=1, gamma*=4. The joint audit verifies that damping and drive interact strongly; it does not define a continuous three-parameter phase diagram.',
       'Recommended Manuscript Wording':'“The sampled vector-locking map is a declared-protocol result; a compact gamma-by-drive audit shows that damping can reorganize the winding map rather than merely shift one plateau boundary.”',
       'Evidence Artifact':'janus_n17_upstream_strengthening/fr_m02_joint_gamma_drive_map.csv',
       'Manuscript Action':'UPDATED after Five Reviewer FR-M02; keep gamma=4 headline, add joint-interaction limitation/evidence in Discussion/SI.'})

c17=df[df.ID=='C17'].iloc[0]
upd('C17',
    **{'Quantitative / Falsification Contract':str(c17['Quantitative / Falsification Contract'])+' Trajectory-count convergence: 500 without-replacement subsamples from pooled N=1000; >=95% Hotelling zero-vector rejection stability occurs by N=100,200,200,300,500 for T*=0.50,0.55,0.60,0.65,0.70.',
       'Scope':str(c17['Scope'])+' The N=1000 high-T design is supported by an explicit precision/convergence audit; this is not a retrospective power calculation.',
       'Evidence Artifact':'janus_n17_upstream_strengthening/fr_s01_thermal_n_convergence.csv',
       'Manuscript Action':'UPDATED after Five Reviewer FR-S01; add N/precision convergence statement to Results/SI.'})

csv_out=OUT/'CLAIM_EVIDENCE_MATRIX_v3.csv';df.to_csv(csv_out,index=False)
# simple markdown source of truth
md=['# Claim-Evidence Matrix v3','',df.to_markdown(index=False)]
(OUT/'CLAIM_EVIDENCE_MATRIX_v3.md').write_text('\n'.join(md),encoding='utf-8')
# update workbook conservatively
xlsx_in=Path('/mnt/data/janus_n13_claim_figure_revision/Janus_MoSSe_ClaimEvidenceMatrix_v2.xlsx');xlsx_out=OUT/'Janus_MoSSe_ClaimEvidenceMatrix_v3.xlsx';shutil.copy2(xlsx_in,xlsx_out)
wb=load_workbook(xlsx_out);ws=wb.active
headers={c.value:i+1 for i,c in enumerate(ws[1])}
for rid in ['M0','C13','C17']:
    row=None
    for r in range(2,ws.max_row+1):
        if ws.cell(r,headers.get('ID',1)).value==rid: row=r;break
    if not row: continue
    src=df[df.ID==rid].iloc[0]
    for col in df.columns:
        if col in headers: ws.cell(row,headers[col]).value=src[col]
wb.save(xlsx_out)

# Updated issue register / gate summary
issues=pd.read_csv(OUT/'N16_FIVE_REVIEWER_ACTION_REGISTER.csv')
issues.loc[issues.issue_id=='FR-M01','status']='CLOSED_WITH_SCOPE'
issues.loc[issues.issue_id=='FR-M01','resolution_path']='Published-decimal precision propagated; sign/winding robust. Actual source fit covariance is unavailable and explicitly retained as NOT_ASSESSABLE input limitation; claims conditional on published Fourier model.'
issues.loc[issues.issue_id=='FR-M02','status']='CLOSED'
issues.loc[issues.issue_id=='FR-M02','resolution_path']='72-point gamma x drive map completed; confirms strong damping-drive interaction and validates protocol-specific wording.'
issues.loc[issues.issue_id=='FR-S01','status']='CLOSED'
issues.loc[issues.issue_id=='FR-S01','resolution_path']='Trajectory-count subsampling convergence completed; high-T N=1000 design is conservative relative to >=95% decision-stability thresholds.'
issues.loc[issues.issue_id=='FR-R01','status']='OPEN_VENUE'
issues.to_csv(OUT/'N17_ISSUE_REGISTER.csv',index=False)
summary={'gate':'N17_UPSTREAM_STRENGTHENING_AND_TARGETED_REVISION','FR-M01':'CLOSED_WITH_SCOPE','FR-M02':'CLOSED','FR-S01':'CLOSED','FR-R01':'OPEN_VENUE_MINOR','scientific_blockers':0,'scientific_major_open':0,'scientific_moderate_open':0,'venue_minor_open':1,'decision':'SCIENTIFIC_STOP_GATE_READY_PENDING_QA','next':'render DOCX, run affected writing QA/closure, then ACS Nano venue handoff'}
(OUT/'N17_GATE_SUMMARY.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
print(main_out);print(si_out);print(csv_out)
