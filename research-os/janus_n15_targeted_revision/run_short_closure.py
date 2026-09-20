from docx import Document
import csv, json, re
from pathlib import Path

root=Path('/mnt/data/janus_n15_targeted_revision')
main=Document(root/'Janus_MoSSe_RedTeamScoped_Manuscript_v4.docx')
si=Document(root/'Janus_MoSSe_RedTeamScoped_SI_v4.docx')
main_text='\n'.join(p.text for p in main.paragraphs)
si_text='\n'.join(p.text for p in si.paragraphs)
all_text=main_text+'\n'+si_text

checks={}
forbidden=[
    'earlier 0.60/0.65/0.70 detectability boundaries',
    'earlier 10-burn + 10-measure pilot boundaries',
    'stationarity audit removes the earlier finite-window thermal boundary',
    'resolved through T*=0.60',
    'unresolved at T*=0.65',
    'resolved at T*=0.65 and unresolved at T*=0.70',
]
checks['legacy_thermal_forbidden_absent']=all(x.lower() not in all_text.lower() for x in forbidden)
checks['legacy_hits']=[x for x in forbidden if x.lower() in all_text.lower()]

required={
    'damping_protocol_specific':'The sampled vector-locking map is therefore a property of the declared m*=1, gamma*=4 reduced dynamical protocol, not a material-only MoSSe phase diagram.',
    'damping_sweep':'windings (1,-10), (1,-3) and (2,-2) at gamma*=1, 2 and 3, retains (1,-1) for gamma*=4-6, and becomes pinned (0,0) for gamma*=8-10',
    'loading_axis':'The headline positive/negative-y comparison is therefore one chosen crystallographic loading axis, not a rotationally invariant friction scalar.',
    'preparation_sign':'the higher-energy minimum loses stability at 0.211059 and 0.611697, giving Delta F_c=-0.400638',
    'boundary_no_kinetics':'No A-to-B transition path, activation barrier, continuous twist-sweep switching protocol, or hysteresis was calculated',
    'mechanics_scope':'spatially uniform center-of-mass/body-force tilt applied at fixed imposed twist',
    'elasticity_scope':'two distinct internal representations that share the same GSFE and target elastic constants',
}
checks['required_scope_present']={k:(v in main_text) for k,v in required.items()}

# Figure 6 source data
figa=list(csv.DictReader(open('/mnt/data/janus_n13_claim_figure_revision/Figure_6_panel_a_data.csv',encoding='utf-8-sig')))
figc=list(csv.DictReader(open('/mnt/data/janus_n13_claim_figure_revision/Figure_6_panel_c_data.csv',encoding='utf-8-sig')))
# inspect headers + row at T=.70
checks['figure6_a_headers']=list(figa[0].keys()) if figa else []
checks['figure6_c_headers']=list(figc[0].keys()) if figc else []
checks['figure6_a_T070']=[r for r in figa if r.get('T')=='0.7' or r.get('T*')=='0.7' or r.get('T')=='0.70' or r.get('T*')=='0.70']
checks['figure6_c_T070']=[r for r in figc if r.get('T')=='0.7' or r.get('T*')=='0.7' or r.get('T')=='0.70' or r.get('T*')=='0.70']

# Extract thermal Table 3 row
thermal_cells=[]
for ti,t in enumerate(main.tables):
    for row in t.rows:
        vals=[c.text.strip() for c in row.cells]
        if vals and vals[0]=='Thermal current':
            thermal_cells=vals
            checks['table3_index']=ti
checks['table3_thermal_row']=thermal_cells

# SI S8b row at 0.70
s8b=[]
for ti,t in enumerate(si.tables):
    headers=[c.text.strip() for c in t.rows[0].cells] if t.rows else []
    if headers and headers[0]=='T*' and 'zero vector' in headers:
        for row in t.rows[1:]:
            vals=[c.text.strip() for c in row.cells]
            if vals and vals[0]=='0.70':
                s8b=vals
                checks['si_s8b_table_index']=ti
checks['si_s8b_T070']=s8b

# Claim matrix C15-C19 rows
mrows={}
with open('/mnt/data/janus_n13_claim_figure_revision/CLAIM_EVIDENCE_MATRIX_v2.csv',encoding='utf-8-sig') as f:
    for r in csv.DictReader(f):
        if r['ID'] in {'C15','C16','C17','C18','C19'}:
            mrows[r['ID']]=r
checks['matrix_C15_C19_present']=sorted(mrows)
checks['matrix_C17_claim']=mrows.get('C17',{}).get('Claim','')
checks['matrix_C18_contract']=mrows.get('C18',{}).get('Quantitative / Falsification Contract','')

# Key synchronization numeric strings
sync_strings=['0.03700','-0.07645','0.01265','0.06132','-0.10088','-0.05195','0.50802','0.01116']
checks['sync_strings_main']={s:(s in main_text) for s in sync_strings}
checks['sync_strings_si']={s:(s in si_text) for s in sync_strings if s not in {'0.50802','0.01116'}}
checks['all_scope_checks_pass']=checks['legacy_thermal_forbidden_absent'] and all(checks['required_scope_present'].values())

(root/'N15_SHORT_CLOSURE_CHECKS.json').write_text(json.dumps(checks,indent=2),encoding='utf-8')
print(json.dumps(checks,indent=2))
