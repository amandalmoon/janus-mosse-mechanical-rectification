from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension
from pathlib import Path
import csv, json, textwrap, zipfile

OUT = Path('/mnt/data/janus_n9_claim_matrix')

claims = [
{
'id':'M0','category':'Material input','claim':'The conservative input is the corrected published three-shell 2H MoSSe Se-S-Se-S GSFE, evaluated with the 2024 Erratum phase convention.','evidence':'VERIFIED','novelty':'NOT_A_NOVELTY_CLAIM','contract':'W=(7.9,0.1,0.1) meV; phi=(131.9°,1.2°,85.4°); independent complex-vs-paired implementation agreement at ~1e-13 or better.','scope':'Published Fourier parameterization; does not independently reproduce the raw 9x9 DFT registry grid.','wording':'Use: “three-shell GSFE coefficients reported by Angeli et al. (2022), evaluated with the Fourier phase convention corrected in the 2024 Erratum.”','forbidden':'Do not call the complete dynamics first-principles or imply the raw DFT grid was independently reproduced.','evidence_artifact':'janus_audit/angeli_differential_audit.json','source_url':'https://doi.org/10.1103/PhysRevB.106.235159 ; https://doi.org/10.1103/PhysRevB.109.199902','manuscript_action':'KEEP; standardize source wording.'
},
{
'id':'C1','category':'Material symmetry','claim':'The corrected 2H MoSSe GSFE has non-removable registry-space inversion asymmetry.','evidence':'VERIFIED','novelty':'NOT_A_NOVELTY_CLAIM','contract':'A_min=0.04253608528083; odd RMS=0.20624278237270; multistart spread ~8e-17; triad invariant chi1=0.583541211356 with translation error <1e-15.','scope':'Conservative lateral registry landscape only.','wording':'“Translation-minimized optimization shows that the corrected 2H MoSSe GSFE is genuinely inversion asymmetric in registry space.”','forbidden':'Do not equate lateral registry inversion asymmetry with experimental out-of-plane polarity reversal.','evidence_artifact':'janus_n4_gate/N4_ODD_SECTOR_GAUGE_GATE_REPORT.md','source_url':'https://doi.org/10.1103/PhysRevB.106.235159 ; https://doi.org/10.1103/PhysRevB.109.199902','manuscript_action':'KEEP; optionally clarify chi1 as a reciprocal-triad product invariant.'
},
{
'id':'C2','category':'Prepared-state depinning','claim':'The N=127 unguided contact has direction-split full-2D prepared-state depinning when loading begins from the lowest-energy zero-force minimum.','evidence':'VERIFIED','novelty':'LIKELY_COMBINATION_ONLY','contract':'At theta=1.5°: Fc+=3.0711353385, Fc-=1.2586722985, rho=0.418601. 491-point continuation over 0-3° finds no ground/meta exchange; all energy and threshold margins remain positive.','scope':'Rigid compact N=127 finite-contact model; prepared branch; k_perp=0.','wording':'“The lowest-energy zero-force registry family exhibits unequal branch-followed loss-of-metastability forces under ±y loading in the unrestricted two-dimensional contact.”','forbidden':'Do not define the physical threshold as the last stable minimum anywhere in the registry cell. Do not call every boundary point a generic saddle-node.','evidence_artifact':'janus_branch_gate/N1_CONTINUOUS_INTERVAL_VERIFICATION_REPORT.md','source_url':'','manuscript_action':'REVISE Methods/Results from 0.25° sampled wording to dense/adaptive numerical verification if strongest claim is desired.'
},
{
'id':'C3','category':'Bifurcation classification','claim':'The prepared depinning boundary is generically fold-like but contains an isolated higher-order stability degeneracy near theta=2.97338566°.','evidence':'VERIFIED_WITH_SCOPE','novelty':'NOT_A_NOVELTY_CLAIM','contract':'Special point: theta≈2.9733856558°, F+≈1.9100735315; one soft mode, positive hard mode ~6.9138; cubic/transversality generic condition fails while metastability is lost across the point.','scope':'N=127 rigid prepared branch.','wording':'Use “depinning/stability boundary” globally; reserve “fold” or “saddle-node” for generic segments.','forbidden':'Do not state that every threshold point is a generic saddle-node.','evidence_artifact':'janus_branch_gate/nongeneric_point_audit.json','source_url':'','manuscript_action':'KEEP current cautious terminology.'
},
{
'id':'C4','category':'Symmetry causality','claim':'Within the reduced same-spectrum model, registry inversion asymmetry controls the directional static and dynamical response.','evidence':'VERIFIED_WITHIN_MODEL','novelty':'LIKELY_AS_INTERNAL_FALSIFICATION','contract':'Matched symmetrized same-2H: global Fc+=Fc- to ~1e-13 and deterministic current ~0. Exact inversion swaps Fc± to ~1e-14 or better and maps (1,-1)→(-1,1).','scope':'Mathematical controls built from the same corrected 2H spectrum; reduced finite-contact model.','wording':'“Same-spectrum symmetrization removes the global directional bias, while exact spatial inversion swaps the two thresholds and reverses the vector current.”','forbidden':'Do not call exact inversion a physical polarization reversal or claim a Pz→response law.','evidence_artifact':'janus_symmetry_gate/N2_SAME2H_SYMMETRY_CAUSALITY_GATE_REPORT.md','source_url':'','manuscript_action':'KEEP as one of the central causal claims.'
},
{
'id':'C5','category':'Canonical model','claim':'The canonical headline baseline is unguided (k_perp=0) with preparation from the lowest-energy zero-force stable minimum; transverse confinement is secondary pathway control.','evidence':'VERIFIED','novelty':'NOT_A_NOVELTY_CLAIM','contract':'Baseline validator PASS; vector, thermal, size/shape, and boundary calculations inherit the unguided prepared-state contract. Legacy production.json with k_perp=25 separated as guided provenance.','scope':'Project/reproducibility definition.','wording':'“Intrinsic” means no added transverse guide in the reduced GSFE model, not a parameter-free physical device.','forbidden':'Do not use guided k_perp=25 as the headline baseline or imply the guide is required for rectification.','evidence_artifact':'janus_baseline_gate/CANONICAL_BASELINE_REPORT.md','source_url':'','manuscript_action':'KEEP; ensure config naming in the final release matches the text.'
},
{
'id':'C6','category':'Finite-size similarity','claim':'For tested compact self-similar rigid contacts, prepared-state threshold surfaces are governed primarily by Theta=theta*sqrt(N), consistent with inverse-linear-size similarity.','evidence':'VERIFIED_WITH_SCOPE','novelty':'NOT_NOVEL_AS_EXPONENT','contract':'Original 10 hexagons: max cross-size deviation 0.2364%/+ and 0.2305%/- for Theta<=18; exact characteristic-angle fits alpha≈0.5036 with R2>=0.99999293; excluding small N drives alpha to ~0.5013.','scope':'Compact, self-similar, rigid contacts; tested sizes/shapes only.','wording':'“theta_q ∝ N^{-1/2} with small finite-size corrections.”','forbidden':'Do not present 0.5036 as a new universal critical exponent.','evidence_artifact':'janus_n2_gate/N2_FINITE_SIZE_SHAPE_GATE_REPORT.md','source_url':'https://doi.org/10.1103/PhysRevB.94.045401 ; https://doi.org/10.1016/j.jmps.2022.105114','manuscript_action':'KEEP; strengthen with exact-root and small-N-exclusion wording if useful.'
},
{
'id':'C7','category':'Compact shape robustness','claim':'The Theta scaling is not specific to one compact perimeter: larger hexagonal and disk-like families nearly coincide through Theta=20.','evidence':'VERIFIED_WITH_SCOPE','novelty':'LIKELY_AS_ROBUSTNESS_TEST','contract':'Hex within-family deviations <=0.0818%/+ and 0.0800%/-; disk <=0.0230%/+ and 0.0222%/-; family-mean difference ~0.105% or less at Theta=20.','scope':'Tested compact rigid hexagon/disk families; not arbitrary edges or deformable flakes.','wording':'“Independent compact shape families agree at approximately the 0.1% level through Theta=20.”','forbidden':'Do not say shape-independent for arbitrary contacts.','evidence_artifact':'janus_n2_gate/n2_extended_shape_dense_metrics.csv','source_url':'https://doi.org/10.1016/j.jmps.2024.105555','manuscript_action':'KEEP.'
},
{
'id':'C8','category':'Boundary registry competition','claim':'The nominal high-twist sign reversal for 20° and 25° triangular cuts is caused by a zero-force ground-state switch between two stable registry families carrying opposite branch-followed biases.','evidence':'VERIFIED_WITH_SCOPE','novelty':'UNCERTAIN','contract':'Prepared switch: 2.921278008381° (20°) and 2.847233122262° (25°); branch splits at crossing approximately +0.665/-0.669 and +0.700/-0.703. Smooth global-envelope zeros occur ~0.008° earlier and are different observables.','scope':'Rigid fixed-N=127 triangular boundary family.','wording':'“Boundary rotation brings competing registry minima into ground-state competition; their energy crossing can reverse the prepared response.”','forbidden':'Do not describe the physical reversal as a smooth zero of the global last-surviving-minimum envelope.','evidence_artifact':'janus_n3_gate/N3_BOUNDARY_REGISTRY_COMPETITION_GATE_REPORT.md','source_url':'','manuscript_action':'KEEP mechanism; remove any legacy global-envelope reversal angles.'
},
{
'id':'C9','category':'Boundary sensitivity','claim':'Registry competition persists under the tested outer-site weighting perturbation, but the switch angle and whether it lies inside 0-3° are strongly edge-model dependent.','evidence':'VERIFIED_WITH_SCOPE','novelty':'NOT_A_NOVELTY_CLAIM','contract':'20° switch range 2.7779-3.1701°; 25° range 2.7108-3.0832°. Only 3/5 and 4/5 tested weights, respectively, switch by 3°.','scope':'Surrogate outer-site local-GSFE weighting 0.5-1.5; not atomistic edge reconstruction.','wording':'“The A/B competition survives, whereas the quantitative switch location—and inclusion in a finite twist window—is edge-model sensitive.”','forbidden':'Do not say reversal necessarily occurs below 3° for all plausible edges.','evidence_artifact':'janus_n3_gate/n3_edge_weight_sensitivity.csv','source_url':'','manuscript_action':'REVISE Discussion 4.2 phrase “switch survives” to “registry competition/crossing persists”; clarify finite-window sensitivity.'
},
{
'id':'C10','category':'Odd-sector mechanism','claim':'The arbitrary-origin slope d(DeltaFc)/deta (including the old ~2.8322 value) is a material susceptibility.','evidence':'REJECTED','novelty':'REJECTED','contract':'Physical eta=1 thresholds are gauge invariant to ~1e-14, but origin-fixed eta slope changes with registry origin; canonical A_min-centered prepared observable has a finite branch-selection jump at eta=0, so the derivative does not exist there.','scope':'Applies to the computational continuation U_eta=U_even+eta U_odd.','wording':'If eta is retained, label it a gauge-fixed or A_min-centered mathematical homotopy for mechanism isolation.','forbidden':'Do not call dDeltaFc/deta a material susceptibility.','evidence_artifact':'janus_n4_gate/N4_ODD_SECTOR_GAUGE_GATE_REPORT.md','source_url':'','manuscript_action':'REMOVE from Abstract/headline Results/Conclusion; SI-only if retained.'
},
{
'id':'C11','category':'Odd-sector mechanism','claim':'The old intermediate optimum eta≈0.5 is a physical/material optimum.','evidence':'REJECTED','novelty':'REJECTED','contract':'Unguided published-origin path peaks near eta≈0.925; A_min-centered path increases to eta=1 over tested interval; previous guided optimum also moves with origin.','scope':'Computational eta homotopy.','wording':'No material optimum should be reported.','forbidden':'Do not report eta_opt≈0.5 as a physical design optimum.','evidence_artifact':'janus_n4_gate/N4_ODD_SECTOR_GAUGE_GATE_REPORT.md','source_url':'','manuscript_action':'REMOVE wherever present in legacy material.'
},
{
'id':'C12','category':'Odd-sector mechanism','claim':'An A_min-centered eta path can be retained as a reproducible mechanism-isolation homotopy.','evidence':'VERIFIED_AS_DIAGNOSTIC','novelty':'NOT_A_NOVELTY_CLAIM','contract':'Transported center under 20 random reparameterizations agrees to ~1e-9; four symmetry-equivalent A_min centers give eta=0.5 thresholds equal to few-1e-8.','scope':'Mathematical diagnostic only; not Pz, not experimental control.','wording':'“A translation-minimized-centered eta homotopy is used only as an SI mechanism-isolation diagnostic.”','forbidden':'Do not map eta directly to physical polarization.','evidence_artifact':'janus_n4_gate/N4_ODD_SECTOR_GAUGE_GATE_REPORT.md','source_url':'','manuscript_action':'OPTIONAL SI ONLY.'
},
{
'id':'C13','category':'Deterministic dynamics','claim':'Zero-mean forcing produces a broad sampled family of integer vector-locked states on the tested 91-point amplitude-period grid.','evidence':'VERIFIED_WITH_SCOPE','novelty':'LIKELY_COMBINATION_ONLY','contract':'91/91 fresh replay points integer-locked; 91/91 also repeat the same integer winding on every measured cycle; max single-cycle locking residual ~9.55e-10.','scope':'theta=1.5°, F0*=1-4, tau*=20-80 sampled grid; m*=1, gamma*=4; deterministic reduced model.','wording':'“All 91 sampled states repeat a single integer lattice-vector winding cycle-by-cycle after transients.”','forbidden':'Do not claim every point in the continuous (F0,tau) plane is locked.','evidence_artifact':'janus_n5_gate/N5_VECTOR_MODE_LOCKING_FLOQUET_GATE_REPORT.md','source_url':'https://doi.org/10.1103/RevModPhys.81.387 ; https://doi.org/10.1103/PhysRevE.62.1988','manuscript_action':'STRENGTHEN Results 3.6 from mean-only integer locking to cycle-resolved 91/91.'
},
{
'id':'C14','category':'Floquet dynamics','claim':'The canonical (1,-1) orbit and representative additional plateaus are attracting relative-periodic orbits.','evidence':'VERIFIED_WITH_SCOPE','novelty':'LIKELY_COMBINATION_ONLY','contract':'Canonical F0*=2,tau*=40: closure ~2e-11, rho_F≈6.31e-22 with DOP853/Radau rho_F<1. Six representative plateau states also rho_F<1; 24/24 timestep and 162/162 local basin probes pass.','scope':'Canonical orbit plus selected representative plateaus; not all 91 states Floquet-certified.','wording':'“Representative plateaus satisfy relative-periodic closure and cross-solver Floquet attraction.”','forbidden':'Do not claim all 91 points are independently Floquet-certified or interpret 1e-22 as a measurable physical scale.','evidence_artifact':'janus_n5_gate/n5_newton_shooting_floquet.csv','source_url':'https://doi.org/10.1103/PhysRevE.62.1988','manuscript_action':'STRENGTHEN carefully; preserve distinction between sampled locking and Floquet-certified representatives.'
},
{
'id':'C15','category':'Thermal dynamics','claim':'Exact deterministic winding identity is much less robust to thermal noise than the mean directed current.','evidence':'VERIFIED','novelty':'LIKELY_AS_MODEL_RESULT','contract':'Across sampled nonzero T*=0.02-0.70, target (1,-1) cycle probability remains <=~0.081 while directed mean displacement persists to much higher T*.','scope':'BAOAB reduced thermal model; 1000 independent trajectories per temperature, 10 burn + 10 measured cycles.','wording':'“Thermal phase slips rapidly mix exact winding labels while a biased mean vector current persists.”','forbidden':'Do not interpret low target-winding probability as immediate loss of directed transport.','evidence_artifact':'janus_n8_gate/N8_THERMAL_ELASTICITY_SENSITIVITY_GATE_REPORT.md','source_url':'','manuscript_action':'KEEP.'
},
{
'id':'C16','category':'Thermal dynamics','claim':'The lattice-v component is statistically resolved through T*=0.60 and unresolved at T*=0.65.','evidence':'VERIFIED_POINTWISE','novelty':'NOT_A_NOVELTY_CLAIM','contract':'Trajectory-cluster 95% CI for mean v excludes 0 through 0.60; at 0.65 it includes 0.','scope':'One component, pointwise temperature-wise CI; not familywise.','wording':'Call this explicitly the lattice-v component, not the full vector current.','forbidden':'Do not use the v-component boundary as the full 2D current boundary.','evidence_artifact':'janus_n8_gate/n8_thermal_vector_joint_audit.csv','source_url':'','manuscript_action':'REVISE Results 3.7, Fig. 6 caption, Table 3, Discussion 4.3 and Conclusion terminology.'
},
{
'id':'C17','category':'Thermal dynamics','claim':'The full two-component mean lattice current remains statistically resolved at the sampled T*=0.65 point and is unresolved at T*=0.70.','evidence':'VERIFIED_POINTWISE','novelty':'LIKELY_AS_MODEL_RESULT','contract':'100k joint bootstrap: at 0.65 zero vector outside simultaneous 95% region; at 0.70 inside. Independent Mahalanobis bootstrap gives same decision.','scope':'Sampled temperatures; two-component joint pointwise inference; no exact transition temperature and no familywise band over T.','wording':'“A joint two-component bootstrap excludes the zero vector at T*=0.65 and includes it at T*=0.70.”','forbidden':'Do not claim a critical temperature between 0.65 and 0.70.','evidence_artifact':'janus_n8_gate/n8_thermal_vector_boundary_100k.csv','source_url':'','manuscript_action':'MANDATORY UPDATE to Abstract, Results 3.7, Table 3, Discussion 4.3 and Conclusion.'
},
{
'id':'C18','category':'Thermal dynamics','claim':'The cycle-sign statistic P(Delta y<0) is clearly biased through T*=0.50, marginal at 0.55, and unresolved by 0.60.','evidence':'VERIFIED_POINTWISE','novelty':'NOT_A_NOVELTY_CLAIM','contract':'100k cluster bootstrap: 0.50 CI above 0.5; 0.55 lower bound ~0.5001; 0.60 and 0.65 include 0.5.','scope':'Cycle-sign statistic only; trajectory-cluster resampling.','wording':'Report separately from mean vector current.','forbidden':'Do not compress all thermal observables into one ad hoc T50 for the unguided 2D system.','evidence_artifact':'janus_n8_gate/n8_Pneg_boundary_100k.csv','source_url':'','manuscript_action':'KEEP current separation; ensure labels are observable-specific.'
},
{
'id':'C19','category':'Thermal numerics','claim':'The low-noise thermal mode mixing/current shift is not caused by the tested timestep range.','evidence':'VERIFIED_WITH_SCOPE','novelty':'NOT_A_NOVELTY_CLAIM','contract':'Independent 300-trajectory reruns at T*=1e-4,1e-3,5e-3 for dt=0.04,0.02,0.01; max pairwise discrepancies ~1.37 combined SE in u and 1.55 in v.','scope':'Targeted low-noise points and timestep range only.','wording':'“Compatible across dt=0.04, 0.02 and 0.01 in targeted low-noise checks.”','forbidden':'Do not imply complete stochastic integrator convergence at all temperatures.','evidence_artifact':'janus_n8_gate/n8_thermal_dt_independent.csv','source_url':'','manuscript_action':'KEEP as SI/numerical robustness.'
},
{
'id':'C20','category':'In-plane elasticity','claim':'Allowing realistic-order in-plane compliance softens absolute depinning forces but preserves the directional sign and leaves rho comparatively stable.','evidence':'VERIFIED_WITH_SCOPE','novelty':'LIKELY_AS_ROBUSTNESS_TEST','contract':'theta=1.5° nominal FEM: Fc+=2.934302, Fc-=1.190210, rho=0.422860 vs rigid rho=0.418601; C/2, 2C and alternative constants retain Fc+>Fc-; 100C returns within <0.07%.','scope':'N=127 free-edge in-plane membrane; no out-of-plane/edge reconstruction.','wording':'“In-plane relaxation renormalizes the absolute pinning scale but does not remove the prepared-state directional split in the tested stiffness range.”','forbidden':'Do not treat rigid Fc values as material constants.','evidence_artifact':'janus_n8_gate/n8_elastic_linear_tight.csv','source_url':'https://doi.org/10.1039/C8CP00350E ; https://doi.org/10.1021/acs.jpclett.3c01066','manuscript_action':'KEEP; refresh numerical values if using new fresh rerun rather than older rounded values.'
},
{
'id':'C21','category':'Nonlinear elasticity','claim':'The in-plane robustness is not a linear-FEM discretization artifact: an independently reconstructed nonlinear bond-angle VFF gives nearly the same thresholds and rho at matched elastic constants.','evidence':'VERIFIED_WITH_SCOPE','novelty':'LIKELY_AS_ROBUSTNESS_TEST','contract':'Reconstructed VFF theta=1.5°: 2.935400/1.190479, rho=0.422921; differs from fresh FEM by ~0.037%/0.023%/0.014%. theta=3° differences remain <0.13%.','scope':'Reconstructed standard VFF matching the documented 342 bonds, 648 angles and target elastic constants; still in-plane only.','wording':'“Two distinct in-plane representations agree closely at matched stiffness.”','forbidden':'Do not call this full atomistic relaxation.','evidence_artifact':'janus_n8_gate/n8_vff_reconstructed.csv','source_url':'https://doi.org/10.1039/C8CP00350E','manuscript_action':'KEEP; archive reconstructed source with final submission package.'
},
{
'id':'C22','category':'Model interpretation','claim':'The rigid-model depinning thresholds are quantitative material constants for an experimental MoSSe slider.','evidence':'REJECTED','novelty':'REJECTED','contract':'In-plane elasticity changes absolute thresholds by ~4-5% at theta=1.5° and ~16-19% at theta=3° for nominal stiffness.','scope':'Interpretation of absolute reduced-model thresholds.','wording':'Describe thresholds as reduced-model predictions and emphasize robustness of sign/normalized split.','forbidden':'Do not present Fc values as universal material constants.','evidence_artifact':'janus_n8_gate/N8_THERMAL_ELASTICITY_SENSITIVITY_GATE_REPORT.md','source_url':'','manuscript_action':'KEEP current limitations language.'
},
{
'id':'C23','category':'Atomistic validation','claim':'The present work validates the effect in a fully 3D atomistically relaxed finite Janus MoSSe contact.','evidence':'NOT_ESTABLISHED','novelty':'NOT_APPLICABLE','contract':'Current hierarchy includes local vertical registry relaxation in the source GSFE plus rigid, linear in-plane FEM and nonlinear in-plane VFF; public SW+KC candidate fails higher-harmonic quantitative GSFE contract.','scope':'Missing collective buckling/corrugation, atomistic edges, defects, nonlocal interactions.','wording':'“The result is DFT-anchored and robust to tested in-plane compliance, but not yet a full 3D atomistic validation.”','forbidden':'Do not use “fully atomistic validation” or “fully first-principles device simulation.”','evidence_artifact':'janus_n8_gate/N8_THERMAL_ELASTICITY_SENSITIVITY_GATE_REPORT.md','source_url':'','manuscript_action':'KEEP as a central limitation.'
},
{
'id':'C24','category':'Polarity mapping','claim':'Physical reversal of Janus polarization Pz is represented by eta→-eta or exact inversion of the lateral GSFE.','evidence':'NOT_ESTABLISHED','novelty':'NOT_APPLICABLE','contract':'No paired polarity-resolved first-principles GSFE surfaces are available; material operation may alter multiple amplitudes and phases.','scope':'Experimental/material interpretation.','wording':'“A quantitative polarity-to-rectification law requires paired polarity-resolved first-principles surfaces or equivalent atomistic input.”','forbidden':'Do not call exact mathematical inversion a physical polarization reversal.','evidence_artifact':'janus_n4_gate/N4_ODD_SECTOR_GAUGE_GATE_REPORT.md','source_url':'','manuscript_action':'KEEP current Discussion 4.4 limitation.'
},
{
'id':'C25','category':'Experimental calibration','claim':'The reduced temperature, damping, mass and drive parameters can already be converted to Kelvin, Hz and device forces.','evidence':'NOT_ESTABLISHED','novelty':'NOT_APPLICABLE','contract':'m*, gamma*, T* and time are reduced/device parameters not supplied by the GSFE input.','scope':'Connection to experiment.','wording':'“No physical thermal or frequency scale is inferred without independent calibration.”','forbidden':'Do not convert T* to kelvin or tau* to physical frequency without calibration.','evidence_artifact':'janus_n8_gate/N8_THERMAL_ELASTICITY_SENSITIVITY_GATE_REPORT.md','source_url':'','manuscript_action':'KEEP limitation.'
},
{
'id':'N1','category':'Novelty positioning','claim':'The generic existence of a 2D friction diode or rocking-ratchet mode locking is itself the novelty.','evidence':'REJECTED_AS_POSITIONING','novelty':'REJECTED','contract':'Closest prior art already includes friction-diode behavior and established rocking-ratchet/phase-locking phenomena.','scope':'Literature positioning.','wording':'Position novelty in the combined DFT-anchored prepared full-2D/same-spectrum/geometry/vector-locking hierarchy.','forbidden':'Do not claim “first friction diode”, “first rocking ratchet”, or generic first mode locking.','evidence_artifact':'Janus_MoSSe_Final_FiveReviewer_Audit.md','source_url':'https://doi.org/10.1021/acsami.6c00303 ; https://doi.org/10.1103/RevModPhys.81.387 ; https://doi.org/10.1103/PhysRevE.62.1988','manuscript_action':'KEEP conservative Introduction positioning.'
},
{
'id':'N2','category':'Novelty positioning','claim':'The combined hierarchy—corrected MoSSe GSFE + unrestricted prepared-state full-2D depinning + same-spectrum symmetry falsification + compact-contact/boundary geometry + lattice-vector relative-periodic transport—is the defensible contribution.','evidence':'SUPPORTED_BY_INTERNAL_EVIDENCE','novelty':'LIKELY; FINAL_REVERSE_SEARCH_PENDING','contract':'All component numerical gates above pass within stated scopes; prior literature audit did not identify the full combination as an existing single study.','scope':'Novelty is a literature statement, not established by numerical validation alone.','wording':'Use cautious contribution language such as “we combine/establish within this model…” rather than an unqualified first-ever claim.','forbidden':'Do not upgrade LIKELY to VERIFIED novelty until final literature reverse search/citation audit is complete.','evidence_artifact':'Janus_MoSSe_Final_FiveReviewer_Audit.md','source_url':'','manuscript_action':'HOLD as LIKELY; revisit in final Citation/Novelty Audit.'
},
]

actions = [
{'priority':'CRITICAL','location':'Abstract','issue':'Thermal sentence currently says mean vector current resolved through T*=0.60 but not 0.65.','required_change':'Split observables: v-component resolved through 0.60/unresolved at 0.65; joint 2D mean vector still resolved at 0.65 and unresolved at 0.70.','claim_ids':'C16,C17'},
{'priority':'MAJOR','location':'Methods 2.3','issue':'Triad diagnostic sin(3 phi1) can be read as a single-phase invariant.','required_change':'Clarify it is the normalized imaginary part/phase of the first-shell complex-amplitude triad product, reducing to sin(3 phi1) in the equal-phase representation.','claim_ids':'C1'},
{'priority':'MAJOR','location':'Methods 2.4 / Results 3.2','issue':'Current text describes the branch audit mainly as a 0.25° grid.','required_change':'Optionally update to the 491-point dense/adaptive interval verification over 0-3°, while retaining the caveat that this is numerical verification rather than interval-arithmetic proof.','claim_ids':'C2,C3'},
{'priority':'MAJOR','location':'Results 3.3','issue':'Current extended shape wording mentions only Theta=0,10,15,20 despite a later dense audit.','required_change':'If desired, state that the larger hex/disk families were subsequently checked on a dense Theta=0..20 grid and retained the same <0.1% conclusion.','claim_ids':'C6,C7'},
{'priority':'MAJOR','location':'Discussion 4.2','issue':'“The switch survives a broad outer-site weighting sensitivity test” can imply the reversal remains within 0-3° for every weight.','required_change':'Replace with: the A/B registry competition and crossing persist across tested edge weights, but the switch angle—and whether it falls inside 0-3°—is edge-model dependent.','claim_ids':'C8,C9'},
{'priority':'MAJOR','location':'Results 3.6','issue':'91-point evidence is described as finite-run mean integer displacement only.','required_change':'Strengthen carefully: 91/91 sampled points repeat the same integer winding cycle-by-cycle after transients; Floquet certification remains limited to the canonical orbit and selected representative plateaus.','claim_ids':'C13,C14'},
{'priority':'CRITICAL','location':'Results 3.7','issue':'Thermal interpretation conflates lattice-v with full vector current.','required_change':'Use observable-specific wording and add the joint two-component 0.65/0.70 result; keep cycle-sign statistic separate.','claim_ids':'C15,C16,C17,C18'},
{'priority':'CRITICAL','location':'Figure 6 caption / Table 3','issue':'Current thermal summary labels the v-component boundary as the mean-vector-current boundary.','required_change':'Label panel/table row explicitly as mean v; add or replace with a full-vector joint-bootstrap line showing zero-vector excluded at 0.65 and included at 0.70.','claim_ids':'C16,C17'},
{'priority':'CRITICAL','location':'Discussion 4.3','issue':'States mean vector current remains nonzero through T*=0.60, which is now incomplete.','required_change':'Distinguish exact winding, cycle-sign, v-component, and joint 2D current. Do not infer a sharp thermal critical temperature.','claim_ids':'C15,C16,C17,C18'},
{'priority':'CRITICAL','location':'Conclusion','issue':'Current conclusion repeats “mean vector current remains resolved through 0.60 but not 0.65.”','required_change':'Replace with the new joint-vector statement and preserve the no-critical-temperature caveat.','claim_ids':'C17'},
{'priority':'MAJOR','location':'Data and code availability / SI','issue':'Standalone nonlinear VFF source was absent from the accessible bundle, although the model was reconstructable.','required_change':'Archive the independently reconstructed VFF source and the new N8 tables in the final reproducibility package.','claim_ids':'C21'},
{'priority':'MAJOR','location':'Legacy materials','issue':'Old eta susceptibility ~2.8322 and eta_opt~0.5 claims remain scientifically invalid if reintroduced.','required_change':'Keep arbitrary-origin eta material-law claims retired. If eta appears, use only A_min-centered SI homotopy language.','claim_ids':'C10,C11,C12'},
]

open_claims = [
{'topic':'Full 3D relaxed contact','status':'OPEN','needed_evidence':'Janus-specific atomistic/MLIP model that reproduces corrected 2H GSFE including higher harmonics, followed by finite-flake relaxation with out-of-plane and edge DOF.','blocks':'Any claim of fully atomistic validation or quantitative experimental threshold prediction.'},
{'topic':'Polarity-response mapping','status':'OPEN','needed_evidence':'Paired polarity-resolved first-principles GSFE surfaces under matched load/separation.','blocks':'Pz→threshold/current transfer law and physical interpretation of eta/inversion as polarity reversal.'},
{'topic':'Physical time/temperature calibration','status':'OPEN','needed_evidence':'Independent effective mass, damping/friction kernel, load and thermal calibration tied to a specific device/contact.','blocks':'Kelvin/Hz/device-force predictions.'},
{'topic':'Boundary switch under realistic edges','status':'OPEN','needed_evidence':'Relaxed/chemistry-specific edge model or atomistic finite flake.','blocks':'Universal or quantitatively material-specific switch angle.'},
{'topic':'Continuous dynamical phase diagram','status':'OPEN','needed_evidence':'2D continuation of relative-periodic branches/tongues and multistability over continuous F0,tau (and possibly theta,k_perp).','blocks':'Claims that the entire continuous parameter plane is locked or globally single-attractor.'},
{'topic':'Final novelty verification','status':'OPEN','needed_evidence':'Final reverse literature search immediately before submission, centered on the combined full-2D prepared-state + same-spectrum + finite-contact + vector-locking hierarchy.','blocks':'Unqualified first-ever novelty wording.'},
]

# Write CSV
csv_path = OUT/'CLAIM_EVIDENCE_MATRIX.csv'
cols = ['id','category','claim','evidence','novelty','contract','scope','wording','forbidden','evidence_artifact','source_url','manuscript_action']
with csv_path.open('w', newline='', encoding='utf-8-sig') as f:
    w=csv.DictWriter(f, fieldnames=cols)
    w.writeheader(); w.writerows(claims)

# JSON
json_path=OUT/'CLAIM_EVIDENCE_MATRIX.json'
json_path.write_text(json.dumps({'claims':claims,'manuscript_actions':actions,'open_claims':open_claims}, ensure_ascii=False, indent=2), encoding='utf-8')

# Markdown
md=[]
md.append('# Janus MoSSe Consolidated Claim–Evidence Matrix — Gate 9')
md.append('')
md.append('## Decision')
md.append('**PASS as a claim-freeze gate, with mandatory manuscript corrections before rewriting.**')
md.append('')
md.append('The numerical story is now internally coherent around one canonical baseline: corrected published 2H MoSSe GSFE → unguided prepared-ground-state full-2D depinning → same-spectrum symmetry falsification → compact-contact / boundary geometry → deterministic vector locking → thermal and in-plane-compliance robustness. The matrix below fixes which statements are verified, scope-limited, rejected, or still open.')
md.append('')
md.append('## Claim matrix')
md.append('')
md.append('| ID | Claim (compressed) | Evidence status | Novelty status | Manuscript action |')
md.append('|---|---|---|---|---|')
for c in claims:
    short=c['claim'].replace('|','/')
    md.append(f"| {c['id']} | {short} | **{c['evidence']}** | {c['novelty']} | {c['manuscript_action']} |")
md.append('')
md.append('## Mandatory manuscript corrections before the rewrite')
md.append('')
for a in actions:
    md.append(f"### {a['priority']} — {a['location']}")
    md.append(f"- Current issue: {a['issue']}")
    md.append(f"- Required change: {a['required_change']}")
    md.append(f"- Claims: {a['claim_ids']}")
    md.append('')
md.append('## Claims that remain explicitly open')
md.append('')
for o in open_claims:
    md.append(f"### {o['topic']} — {o['status']}")
    md.append(f"- Needed evidence: {o['needed_evidence']}")
    md.append(f"- Blocks: {o['blocks']}")
    md.append('')
md.append('## Freeze rules for the manuscript rewrite')
md.append('')
md.append('1. A claim marked VERIFIED may be used only inside its listed scope.')
md.append('2. A claim marked REJECTED must not reappear through legacy text, captions, SI, or figure annotations.')
md.append('3. NOT_ESTABLISHED items belong in limitations/future work, not as positive conclusions.')
md.append('4. Novelty status is separate from numerical validity. LIKELY novelty does not become VERIFIED without a final literature reverse search.')
md.append('5. The new thermal distinction is mandatory: v-component boundary (0.60/0.65) and full-vector boundary (0.65/0.70) are different observables.')
md.append('6. Guided k_perp=25 remains a secondary comparison, never the canonical headline baseline.')
md_path=OUT/'CLAIM_EVIDENCE_MATRIX.md'
md_path.write_text('\n'.join(md), encoding='utf-8')

# Manuscript action register standalone
ar=[]
ar.append('# Manuscript Action Register After Gate 9')
ar.append('')
ar.append('This register contains only changes required by the completed gates; it is not yet the manuscript rewrite.')
ar.append('')
for a in actions:
    ar.append(f"## {a['priority']} — {a['location']}")
    ar.append(f"**Issue:** {a['issue']}")
    ar.append('')
    ar.append(f"**Required change:** {a['required_change']}")
    ar.append('')
    ar.append(f"**Claim IDs:** {a['claim_ids']}")
    ar.append('')
(OUT/'MANUSCRIPT_ACTION_REGISTER.md').write_text('\n'.join(ar), encoding='utf-8')

# Workbook
wb=Workbook()
ws=wb.active; ws.title='Claim-Evidence Matrix'
headers=['ID','Category','Claim','Evidence Status','Novelty Status','Quantitative / Falsification Contract','Scope','Recommended Manuscript Wording','Forbidden / Retired Wording','Evidence Artifact','External Source URL','Manuscript Action']
ws.append(headers)
for c in claims:
    ws.append([c['id'],c['category'],c['claim'],c['evidence'],c['novelty'],c['contract'],c['scope'],c['wording'],c['forbidden'],c['evidence_artifact'],c['source_url'],c['manuscript_action']])

# Styling
header_fill=PatternFill('solid', fgColor='1F4E78')
header_font=Font(color='FFFFFF', bold=True)
status_fills={
    'VERIFIED': 'E2F0D9','VERIFIED_WITH_SCOPE':'E2F0D9','VERIFIED_WITHIN_MODEL':'E2F0D9','VERIFIED_AS_DIAGNOSTIC':'E2F0D9','VERIFIED_POINTWISE':'E2F0D9',
    'SUPPORTED_BY_INTERNAL_EVIDENCE':'DDEBF7','REJECTED':'F4CCCC','REJECTED_AS_POSITIONING':'F4CCCC','NOT_ESTABLISHED':'FCE4D6'
}
for cell in ws[1]: cell.fill=header_fill; cell.font=header_font; cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
for row in ws.iter_rows(min_row=2):
    for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)
    status=row[3].value
    if status in status_fills: row[3].fill=PatternFill('solid',fgColor=status_fills[status])
    novelty=row[4].value or ''
    if 'REJECTED' in novelty: row[4].fill=PatternFill('solid',fgColor='F4CCCC')
    elif 'LIKELY' in novelty: row[4].fill=PatternFill('solid',fgColor='DDEBF7')
    elif 'PENDING' in novelty or 'UNCERTAIN' in novelty: row[4].fill=PatternFill('solid',fgColor='FCE4D6')
    if row[11].value and ('MANDATORY' in row[11].value or 'REVISE' in row[11].value or 'REMOVE' in row[11].value):
        row[11].fill=PatternFill('solid',fgColor='FFF2CC')
ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
widths=[10,22,48,24,27,55,38,55,55,45,45,45]
for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
for r in range(2,ws.max_row+1): ws.row_dimensions[r].height=90
ws.row_dimensions[1].height=35

# Action sheet
wa=wb.create_sheet('Manuscript Actions')
ha=['Priority','Location','Current Issue','Required Change','Claim IDs']
wa.append(ha)
for a in actions: wa.append([a['priority'],a['location'],a['issue'],a['required_change'],a['claim_ids']])
for c in wa[1]: c.fill=header_fill; c.font=header_font; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
for row in wa.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
    p=row[0].value
    row[0].fill=PatternFill('solid',fgColor={'CRITICAL':'F4CCCC','MAJOR':'FFF2CC','MINOR':'DDEBF7'}.get(p,'FFFFFF'))
wa.freeze_panes='A2'; wa.auto_filter.ref=wa.dimensions
for i,w in enumerate([14,28,60,70,18],1): wa.column_dimensions[get_column_letter(i)].width=w
for r in range(2,wa.max_row+1): wa.row_dimensions[r].height=72

# Open claims
wo=wb.create_sheet('Open Claims')
ho=['Topic','Status','Needed Evidence','What It Blocks']
wo.append(ho)
for o in open_claims: wo.append([o['topic'],o['status'],o['needed_evidence'],o['blocks']])
for c in wo[1]: c.fill=header_fill; c.font=header_font; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
for row in wo.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
    row[1].fill=PatternFill('solid',fgColor='FCE4D6')
wo.freeze_panes='A2'; wo.auto_filter.ref=wo.dimensions
for i,w in enumerate([30,14,80,70],1): wo.column_dimensions[get_column_letter(i)].width=w
for r in range(2,wo.max_row+1): wo.row_dimensions[r].height=65

# Evidence index
we=wb.create_sheet('Evidence Index')
he=['Gate / Source','Primary Artifact','Purpose']
we.append(he)
eidx=[
('Material audit','janus_audit/angeli_differential_audit.json','Published GSFE implementation equivalence'),
('N1','janus_branch_gate/N1_CONTINUOUS_INTERVAL_VERIFICATION_REPORT.md','Continuous prepared-branch identity / depinning'),
('Same-2H symmetry','janus_symmetry_gate/N2_SAME2H_SYMMETRY_CAUSALITY_GATE_REPORT.md','Symmetrization and exact inversion causal contracts'),
('Canonical baseline','janus_baseline_gate/CANONICAL_BASELINE_REPORT.md','Unguided prepared-state dependency graph'),
('N2','janus_n2_gate/N2_FINITE_SIZE_SHAPE_GATE_REPORT.md','Finite-size / compact-shape similarity'),
('N3','janus_n3_gate/N3_BOUNDARY_REGISTRY_COMPETITION_GATE_REPORT.md','Boundary registry competition and edge sensitivity'),
('N4','janus_n4_gate/N4_ODD_SECTOR_GAUGE_GATE_REPORT.md','Gauge dependence / rejected eta material law'),
('N5','janus_n5_gate/N5_VECTOR_MODE_LOCKING_FLOQUET_GATE_REPORT.md','Vector locking / Floquet / continuation / basin'),
('N8','janus_n8_gate/N8_THERMAL_ELASTICITY_SENSITIVITY_GATE_REPORT.md','Thermal statistics and in-plane compliance'),
('Current manuscript','Janus_MoSSe_WorkflowValidated_Draft.docx','Current prose to be rewritten after claim freeze'),
]
for x in eidx: we.append(x)
for c in we[1]: c.fill=header_fill; c.font=header_font; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
for row in we.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
for i,w in enumerate([26,65,70],1): we.column_dimensions[get_column_letter(i)].width=w
we.freeze_panes='A2'

# Legend
wl=wb.create_sheet('Legend')
legend=[
('VERIFIED','Numerically/source verified within the explicitly stated scope.'),
('VERIFIED_WITH_SCOPE','Verified, but the scope limitation must travel with the claim.'),
('VERIFIED_POINTWISE','Statistical statement verified at sampled points; not a continuous transition or familywise statement.'),
('REJECTED','Falsified or conceptually invalid; must not return to the manuscript.'),
('NOT_ESTABLISHED','Evidence is absent; belongs in limitations/future work.'),
('LIKELY_COMBINATION_ONLY','Numerical claim is supported, but novelty concerns the combined research architecture, not the generic phenomenon.'),
('FINAL_REVERSE_SEARCH_PENDING','Novelty cannot be treated as verified until a final literature reverse search is completed.'),
]
wl.append(['Label','Meaning'])
for x in legend: wl.append(x)
for c in wl[1]: c.fill=header_fill; c.font=header_font
for row in wl.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
wl.column_dimensions['A'].width=32; wl.column_dimensions['B'].width=100

xlsx_path=OUT/'Janus_MoSSe_ClaimEvidenceMatrix.xlsx'
wb.save(xlsx_path)

# Validation reopen
wb2=load_workbook(xlsx_path, data_only=False)
assert wb2['Claim-Evidence Matrix'].max_row == len(claims)+1
assert wb2['Manuscript Actions'].max_row == len(actions)+1
assert wb2['Open Claims'].max_row == len(open_claims)+1
assert not any(cell.value in ('#REF!','#DIV/0!','#VALUE!','#N/A','#NAME?') for wsx in wb2.worksheets for row in wsx.iter_rows() for cell in row)

# Package
zip_path=OUT/'N9_ClaimEvidence_GateBundle.zip'
with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED) as z:
    for p in [csv_path,json_path,md_path,OUT/'MANUSCRIPT_ACTION_REGISTER.md',xlsx_path]:
        z.write(p, arcname=p.name)

print(f'claims={len(claims)} actions={len(actions)} open={len(open_claims)}')
print(xlsx_path)
print(zip_path)
