from docx import Document
from pathlib import Path
from copy import deepcopy
import zipfile, hashlib, shutil, tempfile, os
import pandas as pd

ROOT=Path('/mnt/data/janus_n13_claim_figure_revision')
SRC_MAIN=Path('/mnt/data/janus_n11_citation_novelty/Janus_MoSSe_CitationAudited_Manuscript_v2.docx')
SRC_SI=Path('/mnt/data/janus_n11_citation_novelty/Janus_MoSSe_CitationAudited_SI_v2.docx')
OUT_MAIN=ROOT/'Janus_MoSSe_ThermalStationary_Manuscript_v3.docx'
OUT_SI=ROOT/'Janus_MoSSe_ThermalStationary_SI_v3.docx'
NEW_FIG=ROOT/'Figure_6_thermal_stationary_hierarchy.png'
OLD_FIG=Path('/mnt/data/janus_n10_rewrite/Figure_6_thermal_vector_hierarchy.png')

# ---------- main manuscript ----------
doc=Document(SRC_MAIN)
repl={
3:"Twist suppresses interfacial corrugation in many incommensurate two-dimensional contacts, but a laterally inversion-asymmetric registry landscape can make finite-contact pinning direction dependent. We analyze the published three-shell generalized stacking-fault-energy (GSFE) landscape of 2H Janus MoSSe using an unguided finite contact whose center of mass is free in the full two-dimensional registry plane. Headline depinning thresholds are followed from the lowest-energy zero-force minimum rather than selected from an unconstrained last-surviving-minimum envelope. For the N=127 reference contact at theta=1.5 deg, the prepared rigid branch loses metastability at F_c,+*=3.071135 and F_c,-*=1.258672. Same-spectrum symmetrization removes the global directional bias and exact spatial inversion swaps the thresholds and reverses the deterministic current. Compact hexagonal and disk-like contacts collapse under theta sqrt(N), rotated triangular cuts can reverse the prepared response through registry-state switching, and zero-mean forcing produces cycle-resolved integer lattice-vector transport over all 91 sampled amplitude-period points. Representative plateaus are attracting under independent Floquet checks. Thermal exact-winding identity is fragile, whereas under a stationarity-audited protocol a weaker mean current vector remains statistically nonzero throughout the sampled range to T*=0.70; no thermal extinction scale is inferred. Two in-plane relaxation representations preserve the directional split while renormalizing absolute thresholds. The claims are therefore DFT-anchored reduced-model results, not a parameter-free experimental device prediction.",
37:"integrated with BAOAB. For the unguided system a scalar half-cycle label is insufficient because trajectories move across both lattice directions. The stationarity-audited thermal protocol uses dt=0.02, 60 burn cycles and 100 measured cycles. A full T*=0.02-0.70 series uses N=500 independent trajectories per temperature, and the high-temperature band T*=0.50-0.70 is independently replicated with a second N=500 seed ensemble. All cycle data from a trajectory are reduced first to trajectory-level observables; inference is then performed over trajectories rather than individual cycles. Late 20-cycle blocks are compared by paired two-dimensional tests, deterministic-orbit/ground/metastable initial states are checked under common noise for memory loss, and high-temperature mean currents are assessed with trajectory-level bootstrap and Hotelling statistics. We separately report mean u and v, P(v_cycle<0), and the probability that a cycle is assigned to the deterministic target (m,n)=(1,-1). Targeted dt=0.04, 0.02 and 0.01 simulations test high-temperature effect-size compatibility as well as the previously checked low-noise regime.",
79:"Thermal robustness is quantified from stationary long-window trajectories, with the trajectory rather than the individual cycle used as the inference unit. The full T*=0.02-0.70 series uses N=500 trajectories after 60 burn cycles and 100 measured cycles, and T*=0.50-0.70 is independently replicated with a second N=500 ensemble. Exact deterministic winding identity is fragile: over the sampled nonzero temperatures the cycle-level probability assigned to the target (1,-1) cell never exceeds about 0.0818, and in the pooled high-temperature band it is only about 0.011-0.016. This does not imply loss of directed transport because noisy cycles can populate many lattice-vector channels while retaining a small nonzero mean drift.",
80:"The stationarity audit removes the earlier finite-window thermal boundary. Late measurement blocks show no resolved two-component drift in either high-temperature seed ensemble, and deterministic-orbit, zero-force ground-state and competing-metastable starts lose their displacement memory by cycle 38 at the latest under matched noise. In the long-window N=500 temperature series the mean-current magnitude decreases monotonically with T*. Pooling the two independent high-temperature ensembles gives mean (u,v)=(0.08001,-0.15509) at T*=0.50 and (0.03700,-0.07645) at T*=0.70. At T*=0.70 a 50,000-resample trajectory bootstrap gives u in [0.01265,0.06132] and v in [-0.10088,-0.05195], and the zero-current vector remains outside the 95% joint region. The cycle-sign statistic approaches one-half but remains weakly biased at the largest sampled temperature, P(v_cycle<0)=0.50802 [0.50506,0.51100], while the target-winding probability is only 0.01116. These are finite-protocol effect-size and detectability statements: no current-extinction temperature, equilibrium transition or behavior beyond T*=0.70 is established.",
83:"Figure 6. Stationary thermal hierarchy from trajectory-level inference. (a) Mean lattice-coordinate currents <u> and <v> after 60 burn cycles and 100 measured cycles. The low-temperature full series uses N=500 trajectories per T*, while T*=0.50-0.70 uses the pooled two-seed N=1000 estimates; error bars are trajectory-level 95% intervals. (b) High-temperature Hotelling zero-current rejection ratio for the pooled two-seed ensembles; values above one reject the zero-current vector under the declared test and are not interpreted as a thermal critical scale. (c) The negative-v cycle probability approaches one-half while the exact deterministic target-winding probability remains near 1% in the high-temperature band, separating weak directed drift from preservation of the deterministic mode identity.",
94:"A static threshold inequality does not by itself determine a rocking-ratchet phase diagram. Generic rectification and phase locking on periodic substrates are established [19,20,23,24]; the material-specific question here is whether the same finite MoSSe contact that exhibits prepared-state directional depinning supports lattice-vector relative-periodic transport under zero-mean longitudinal forcing. The 91-point F0-tau scan is stronger than a mean-displacement classification because every sampled state repeats one winding cycle-by-cycle after transients, and representative plateaus spanning several windings are independently verified as attracting relative-periodic states. Thermal trajectories separate exact mode identity from stationary directed transport. After explicit burn-in, block-stationarity, initial-state-memory, independent-seed and timestep checks, the deterministic target winding is already strongly mixed while a smaller mean drift remains statistically nonzero throughout the sampled range to T*=0.70. The negative-v cycle fraction approaches one-half as temperature rises, showing that a weak bias can survive even when individual cycles are nearly balanced and exact mode identity is rare. The earlier 0.60/0.65/0.70 detectability boundaries obtained from short observation windows are therefore retired. No universal thermal critical temperature or complete operating diagram is inferred outside the tested drive, damping and reduced-temperature ranges.",
100:"A finite contact constructed from the published 2H Janus MoSSe GSFE exhibits branch-followed full two-dimensional mechanical rectification without an added transverse guide. At theta=1.5 deg the prepared N=127 rigid thresholds are F_c,+*=3.071135 and F_c,-*=1.258672, and dense/adaptive continuation over 0-3 deg finds no exchange between the prepared ground family and the competing metastable family. Translation-minimized asymmetry diagnostics and matched same-2H controls identify registry-space inversion asymmetry as the causal conservative ingredient within the reduced model. Compact-contact thresholds obey theta sqrt(N) similarity with finite-size corrections toward the inverse-linear-size law, while rotated triangular boundaries can reverse the prepared bias through a ground-state registry switch whose angle is edge-model dependent. Zero-mean forcing generates cycle-resolved integer lattice-vector transport over all 91 sampled amplitude-period points, with representative plateaus independently verified as attracting. Thermal exact-winding identity is fragile, but after stationarity and replication checks a weaker mean vector drift remains statistically nonzero through the largest sampled T*=0.70; no extinction temperature is established. Linear FEM and independently reconstructed nonlinear VFF relaxations soften absolute thresholds while preserving the directional split. The resulting evidence supports a self-consistent DFT-anchored reduced-model mechanism; it does not establish a unique experimental friction coefficient, a physical polarity law, full 3D atomistic relaxation, a Kelvin-scale thermal boundary or a universal nonlinear phase diagram.",
102:"The accompanying numerical materials contain the canonical unguided rigid-model release, the linear elastic-validation package, the independently reconstructed nonlinear bond-angle source and outputs, the Tier-2 interlayer-potential prescreen, and the Gate-level audit records used to freeze manuscript claims. The rigid and audit packages include the published-GSFE engine; dense prepared-branch enumeration and continuation; matched symmetry and exact-inversion controls; registry-origin and odd-sector gauge audits; exact finite-size/shape and boundary analyses; cycle-resolved vector-locking and Floquet checks; stationary thermal trajectories with block-drift, initialization-memory, independent-seed, bootstrap/Hotelling and timestep audits; and elasticity validation outputs. A citation-audit patch removes an earlier source-comment attribution that incorrectly described the 2024 Erratum as a GSFE phase-convention correction; no numerical coefficient or computed result changes under that documentation correction. No archival DOI is claimed until external repository deposition is completed."
}
for i,text in repl.items():
    doc.paragraphs[i].text=text

# Update main tables by row labels.
for table in doc.tables:
    for row in table.rows:
        first=row.cells[0].text.strip()
        if first=='Thermal inference':
            row.cells[1].text='Stationary BAOAB: dt=0.02; 60 burn + 100 measured cycles; N=500 per T* over 0.02-0.70; independent second N=500 seed at T*=0.50-0.70; trajectory-level stationarity/bootstrap/Hotelling checks'
            row.cells[2].text='Stationary current decay, seed replication and finite-observation robustness'
        if first=='Thermal current':
            row.cells[1].text='stationary current nonzero through sampled T*=0.70; at 0.70 pooled <u,v>=(0.03700,-0.07645), bootstrap u [0.01265,0.06132], v [-0.10088,-0.05195]; P target=0.01116'
            row.cells[2].text='Weak directed drift persists after exact mode identity is strongly mixed; no extinction temperature established'

doc.save(OUT_MAIN)

# Replace the exact old Figure 6 media payload while preserving the existing Word extent/cross-reference.
def replace_media_by_hash(docx_path, old_file, new_file):
    old_hash=hashlib.sha256(old_file.read_bytes()).hexdigest()
    tmp=docx_path.with_suffix('.tmp.docx')
    replaced=[]
    with zipfile.ZipFile(docx_path,'r') as zin, zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data=zin.read(item.filename)
            if item.filename.startswith('word/media/') and hashlib.sha256(data).hexdigest()==old_hash:
                data=new_file.read_bytes(); replaced.append(item.filename)
            zout.writestr(item,data)
    os.replace(tmp,docx_path)
    if not replaced:
        raise RuntimeError('Figure 6 media payload not found')
    return replaced
replace_media_by_hash(OUT_MAIN, OLD_FIG, NEW_FIG)

# ---------- Supporting Information ----------
si=Document(SRC_SI)
si.paragraphs[37].text='Table S8a. Stationary free-contact thermal component and directional statistics (60 burn + 100 measured cycles; N=500 per T*).'
si.paragraphs[38].text='Table S8b. High-temperature stationary two-seed pooled joint-current audit (N=1000 per T*).'
si.paragraphs[40].text=(
    'Each trajectory contains 60 burn cycles and 100 measured cycles at dt=0.02. The full T*=0.02-0.70 series uses N=500 independent trajectories per temperature, and the high-temperature band T*=0.50-0.70 is repeated with a second independent N=500 ensemble. The trajectory is the inferential unit. Late 20-cycle blocks show no resolved two-component drift in either high-temperature seed ensemble, and deterministic-orbit, zero-force ground-state and competing-metastable starts become identical to numerical precision by cycle 38 at the latest under common noise. The pooled high-temperature mean vector decreases from (0.08001,-0.15509) at T*=0.50 to (0.03700,-0.07645) at T*=0.70, but a 50,000-resample trajectory bootstrap still excludes the zero vector at every sampled high-temperature point. At T*=0.70 the component intervals are u in [0.01265,0.06132] and v in [-0.10088,-0.05195]. P(v_cycle<0) approaches one-half but remains slightly above it through 0.70, while the exact target-winding probability is only about 1.1-1.6% over T*=0.50-0.70. The earlier 10-burn + 10-measure pilot boundaries shift with observation length and are therefore retired rather than interpreted as physical thermal thresholds. Targeted dt=0.04, 0.02 and 0.01 reruns at T*=0.60, 0.65 and 0.70 agree within 1.12 combined standard errors in either mean-current component.'
)

full=pd.read_csv('/mnt/data/janus_redteam_gate/rtb01/thermal_stationary_series_N500.csv')
boot=pd.read_csv('/mnt/data/janus_redteam_gate/rtb01/stationary_thermal_bootstrap_50k.csv')
pooled=pd.read_csv('/mnt/data/janus_redteam_gate/rtb01/stationary_current_two_seed_pooled.csv')
pooled=pooled[pooled['seed'].astype(str)=='pooled'].copy()
sign=pd.read_csv('/mnt/data/janus_redteam_gate/rtb01/stationary_sign_target_bootstrap.csv')

# Table S8a is table index 8. Keep six columns, overwrite rows.
t=si.tables[8]
headers=['T*','mean u','mean v','95% CI v','P(v_cycle<0)','P target (1,-1)']
for j,h in enumerate(headers): t.cell(0,j).text=h
rows=[]
for _,r in full.iterrows():
    rows.append([f"{r['T']:.2f}",f"{r['mean_u']:.4f}",f"{r['mean_v']:.4f}",f"[{r['v_ci_lo']:.4f},{r['v_ci_hi']:.4f}]",f"{r['P_v_negative_cycle']:.4f}",f"{r['P_target_cycle']:.4f}"])
for i,rowvals in enumerate(rows, start=1):
    for j,v in enumerate(rowvals): t.cell(i,j).text=v

# Table S8b is table index 9. Need 5 high-T rows; clone last data row to preserve style.
t2=si.tables[9]
while len(t2.rows)<6:
    tr=deepcopy(t2.rows[-1]._tr)
    t2._tbl.append(tr)
headers2=['T*','mean u','95% CI u','mean v','95% CI v','T_H^2/T_0.95^2','zero vector']
for j,h in enumerate(headers2): t2.cell(0,j).text=h
merged=boot.merge(pooled[['T','T2','crit']],on='T').sort_values('T')
for i,(_,r) in enumerate(merged.iterrows(),start=1):
    vals=[f"{r['T']:.2f}",f"{r['mean_u']:.5f}",f"[{r['u_ci_lo']:.5f},{r['u_ci_hi']:.5f}]",f"{r['mean_v']:.5f}",f"[{r['v_ci_lo']:.5f},{r['v_ci_hi']:.5f}]",f"{r['T2']/r['crit']:.2f}",'excluded' if r['zero_excluded'] else 'included']
    for j,v in enumerate(vals): t2.cell(i,j).text=v

si.save(OUT_SI)

# Plain-text exports for deterministic diff/search QA.
for path,outtxt in [(OUT_MAIN,ROOT/'Janus_MoSSe_ThermalStationary_Manuscript_v3.txt'),(OUT_SI,ROOT/'Janus_MoSSe_ThermalStationary_SI_v3.txt')]:
    d=Document(path)
    with open(outtxt,'w',encoding='utf-8') as f:
        for i,p in enumerate(d.paragraphs):
            if p.text.strip(): f.write(f'P{i}: {p.text}\n')
        for ti,table in enumerate(d.tables):
            f.write(f'\nTABLE {ti}\n')
            for row in table.rows:
                f.write(' || '.join(c.text.replace('\n',' / ') for c in row.cells)+'\n')

print(OUT_MAIN)
print(OUT_SI)
