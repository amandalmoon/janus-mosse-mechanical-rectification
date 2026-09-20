from openpyxl import load_workbook
from copy import copy
from pathlib import Path
import csv, json

SRC=Path('/mnt/data/janus_n9_claim_matrix/Janus_MoSSe_ClaimEvidenceMatrix.xlsx')
OUT=Path('/mnt/data/janus_n13_claim_figure_revision')
DST=OUT/'Janus_MoSSe_ClaimEvidenceMatrix_v2.xlsx'
wb=load_workbook(SRC)
ws=wb['Claim-Evidence Matrix']
headers=[c.value for c in ws[1]]
idx={ws.cell(r,1).value:r for r in range(2,ws.max_row+1)}

def setrow(cid, vals):
    r=idx[cid]
    for c,h in enumerate(headers,1):
        if h in vals:
            ws.cell(r,c).value=vals[h]

# Provenance correction carried forward from Gate 11.
setrow('M0', {
    'Claim':'The conservative input is the published three-shell 2H MoSSe Se-S-Se-S GSFE reported by Angeli et al. (2022); the real paired Fourier form follows the source C3 symmetry and reality condition. The 2024 Erratum does not report a GSFE phase-convention correction.',
    'Evidence Status':'VERIFIED',
    'Novelty Status':'NOT_A_NOVELTY_CLAIM',
    'Quantitative / Falsification Contract':'W=(7.9,0.1,0.1) meV; phi=(131.9,1.2,85.4) deg; independent complex-vs-paired implementation agreement at ~1e-13 or better. Gate 11 provenance audit verified the Erratum contents separately.',
    'Scope':'Published Fourier parameterization; does not independently reproduce the raw 9x9 DFT registry grid.',
    'Recommended Manuscript Wording':'Use: “three-shell GSFE coefficients reported by Angeli et al. (2022), written in the real conjugate-paired form implied by C3 symmetry and a real scalar energy.”',
    'Forbidden / Retired Wording':'Do not say the 2024 Erratum corrected the GSFE Fourier phase convention. Do not call the complete dynamics first-principles.',
    'Evidence Artifact':'janus_n11_citation_novelty/ANGELI_GSFE_PROVENANCE_CORRECTION.md',
    'Manuscript Action':'ALREADY FIXED in citation-audited v2; keep matrix synchronized.'
})

setrow('C15', {
    'Claim':'Exact deterministic winding identity is much less robust to thermal noise than the stationary directed mean current.',
    'Evidence Status':'VERIFIED_WITH_SCOPE',
    'Novelty Status':'LIKELY_AS_MODEL_RESULT',
    'Quantitative / Falsification Contract':'Stationary protocol: dt=0.02, 60 burn + 100 measured cycles. Full T*=0.02-0.70 series uses N=500 trajectories; T*=0.50-0.70 has a second independent N=500 seed. Target (1,-1) cycle probability never exceeds ~0.0818 over the sampled nonzero-T series and is only ~0.011-0.016 in the pooled high-T band, while the mean current remains nonzero.',
    'Scope':'Canonical reduced BAOAB model at N=127, theta=1.5 deg, F0*=2, tau*=40, m*=1, gamma*=4. No extrapolation beyond T*=0.70.',
    'Recommended Manuscript Wording':'“Thermal noise rapidly destroys exact cycle-by-cycle mode identity, while a weaker directed mean drift persists throughout the sampled stationary range.”',
    'Forbidden / Retired Wording':'Do not interpret low target-winding probability as immediate loss of directed transport or as a thermal phase transition.',
    'Evidence Artifact':'janus_redteam_gate/rtb01/RTB01_THERMAL_STATIONARITY_CLOSURE_REPORT.md',
    'Manuscript Action':'UPDATE Abstract, Results 3.7, Fig. 6, Discussion and Conclusion to the stationary protocol.'
})
setrow('C16', {
    'Claim':'Under the stationary long-window protocol, the mean lattice-v current remains negative and statistically separated from zero through the largest sampled temperature T*=0.70.',
    'Evidence Status':'VERIFIED_WITH_SCOPE',
    'Novelty Status':'NOT_A_NOVELTY_CLAIM',
    'Quantitative / Falsification Contract':'Pooled two-seed N=1000 high-T audit: at T*=0.70, mean v=-0.07645 with 50k trajectory-bootstrap 95% interval [-0.10088,-0.05195]. The earlier 10+10 claim “resolved through 0.60 / unresolved at 0.65” fails the stationarity audit and is retired.',
    'Scope':'Sampled T*=0.50-0.70 high-T band; trajectory is the inferential unit; detectability depends on N and measurement duration.',
    'Recommended Manuscript Wording':'“The stationary mean v component remains negative through the largest sampled T*=0.70, although its magnitude decays strongly with temperature.”',
    'Forbidden / Retired Wording':'Retire: “mean v is unresolved at T*=0.65” and any component-wise critical-temperature interpretation.',
    'Evidence Artifact':'janus_redteam_gate/rtb01/stationary_thermal_bootstrap_50k.csv',
    'Manuscript Action':'REPLACE old component-boundary language in Results/SI/Figure caption.'
})
setrow('C17', {
    'Claim':'The full two-component stationary mean lattice current remains statistically nonzero through the largest sampled temperature T*=0.70.',
    'Evidence Status':'VERIFIED_WITH_SCOPE',
    'Novelty Status':'LIKELY_AS_MODEL_RESULT',
    'Quantitative / Falsification Contract':'At T*=0.70, pooled N=1000 mean (u,v)=(0.036998,-0.076448). 50k trajectory bootstrap gives u [0.01265,0.06132], v [-0.10088,-0.05195]; zero-vector Mahalanobis distance 38.07 versus 95% contour threshold 6.04. Zero is excluded at every sampled high-T point 0.50-0.70.',
    'Scope':'Stationary periodic reduced model under the declared drive/damping. No critical temperature is located and no statement is made for T*>0.70.',
    'Recommended Manuscript Wording':'“The stationary mean vector decreases strongly with T* but remains statistically nonzero through the largest sampled value T*=0.70.”',
    'Forbidden / Retired Wording':'Retire: “resolved at 0.65 and unresolved at 0.70.” Do not infer a critical thermal scale from finite-observation significance.',
    'Evidence Artifact':'janus_redteam_gate/rtb01/stationary_thermal_bootstrap_50k.csv',
    'Manuscript Action':'MANDATORY UPDATE to Abstract, Results 3.7, Fig. 6, Table 3, Discussion, Conclusion and SI.'
})
setrow('C18', {
    'Claim':'The cycle-sign statistic approaches an unbiased 0.5 at high T* but remains weakly above 0.5 through the largest sampled T*=0.70 under the stationary protocol.',
    'Evidence Status':'VERIFIED_WITH_SCOPE',
    'Novelty Status':'NOT_A_NOVELTY_CLAIM',
    'Quantitative / Falsification Contract':'Pooled high-T trajectory bootstrap: P(v_cycle<0)=0.51933 [0.51632,0.52237] at 0.50 and 0.50802 [0.50506,0.51100] at 0.70; target-winding probability simultaneously falls to ~0.011.',
    'Scope':'Effect is statistically resolved but small; its detectability depends on trajectory count/window and should not define a physical boundary.',
    'Recommended Manuscript Wording':'“The negative-v cycle fraction approaches one-half while retaining a small resolved bias through T*=0.70; exact deterministic winding identity is far more strongly mixed.”',
    'Forbidden / Retired Wording':'Retire: “P(Delta y<0) becomes unresolved by T*=0.60.” Do not define an ad hoc T50 for the unguided vector problem.',
    'Evidence Artifact':'janus_redteam_gate/rtb01/stationary_sign_target_bootstrap.csv',
    'Manuscript Action':'UPDATE Figure 6/SI; retain as a distinct observable from the mean vector current.'
})
setrow('C19', {
    'Claim':'The stationary high-temperature mean-current effect size is compatible across the tested BAOAB timestep range.',
    'Evidence Status':'VERIFIED_WITH_SCOPE',
    'Novelty Status':'NOT_A_NOVELTY_CLAIM',
    'Quantitative / Falsification Contract':'At T*=0.60,0.65,0.70, independent dt=0.04,0.02,0.01 reruns differ by no more than 1.12 combined SE in either component. Separate low-noise checks also showed compatibility across the same timestep range.',
    'Scope':'Tested temperatures, trajectory counts and dt range only; not a proof of global stochastic-integrator convergence.',
    'Recommended Manuscript Wording':'“Mean-current estimates are compatible across dt=0.04, 0.02 and 0.01 in targeted low- and high-temperature checks.”',
    'Forbidden / Retired Wording':'Do not imply complete stochastic-integrator convergence over all temperatures, observables or drive parameters.',
    'Evidence Artifact':'janus_redteam_gate/rtb01/highT_dt_pairwise_z.csv',
    'Manuscript Action':'UPDATE SI numerical-robustness statement; no new headline claim.'
})

# Evidence index: keep N8 for elasticity and add RT-B01.
wei=wb['Evidence Index']
for r in range(2, wei.max_row+1):
    if wei.cell(r,1).value=='N8':
        wei.cell(r,3).value='Earlier thermal pilot plus in-plane compliance; thermal boundary claims superseded by RT-B01.'
        break
newr=wei.max_row+1
# copy style from prior row
for c in range(1,4):
    src=wei.cell(newr-1,c); dst=wei.cell(newr,c)
    if src.has_style:
        dst._style=copy(src._style)
    dst.font=copy(src.font); dst.fill=copy(src.fill); dst.border=copy(src.border); dst.alignment=copy(src.alignment); dst.number_format=src.number_format
wei.cell(newr,1).value='RT-B01'
wei.cell(newr,2).value='janus_redteam_gate/rtb01/RTB01_THERMAL_STATIONARITY_CLOSURE_REPORT.md'
wei.cell(newr,3).value='Thermal burn-in, measurement-window, seed, initialization-memory, bootstrap/Hotelling and timestep stationarity closure.'

# Add open thermal-boundary item if absent.
woc=wb['Open Claims']
if not any('Thermal current extinction' in str(woc.cell(r,1).value) for r in range(2,woc.max_row+1)):
    rr=woc.max_row+1
    for c in range(1,woc.max_column+1):
        src=woc.cell(rr-1,c); dst=woc.cell(rr,c)
        if src.has_style: dst._style=copy(src._style)
        dst.font=copy(src.font); dst.fill=copy(src.fill); dst.border=copy(src.border); dst.alignment=copy(src.alignment); dst.number_format=src.number_format
    vals=['Thermal current extinction beyond T*=0.70','OPEN','Extend the converged stationary protocol beyond T*=0.70 with predeclared precision/power targets and distributional diagnostics.','Any claim of a critical reduced temperature or current extinction point.']
    for c,v in enumerate(vals,1): woc.cell(rr,c).value=v

# Manuscript Actions: replace thermal entries with stationary protocol actions.
wma=wb['Manuscript Actions']
updates={
    2:['CRITICAL','Abstract','Thermal sentence retains the superseded 0.65/0.70 detectability boundary.','Replace with stationary claim: exact winding is strongly mixed, while a weaker mean vector drift remains nonzero through the largest sampled T*=0.70; no critical temperature is inferred.','C15,C17'],
    8:['CRITICAL','Results 3.7','Results use the superseded 10+10 component/joint boundaries.','Replace with 60-burn + 100-measure stationary protocol, two-seed high-T replication, monotonic current decay, and nonzero current through T*=0.70.','C15,C16,C17,C18'],
    9:['CRITICAL','Figure 6 caption / Table 3','Figure 6 and Table 3 encode a false 0.65/0.70 thermal boundary.','Regenerate Figure 6 from RT-B01 stationary data and replace Table 3 thermal row with the stationary protocol and T*=0.70 evidence.','C15,C16,C17,C18'],
    10:['CRITICAL','Discussion 4.3','Discussion presents observable-specific detectability thresholds as a thermal hierarchy.','Replace threshold ordering with distinction between exact mode identity, cycle-sign bias and stationary mean drift; state that no extinction temperature is established.','C15,C16,C17,C18'],
    11:['CRITICAL','Conclusion','Conclusion repeats the superseded 0.65/0.70 boundary.','State only that the stationary mean drift persists through the sampled range to T*=0.70 while exact winding identity is fragile.','C15,C17'],
}
for r,vals in updates.items():
    for c,v in enumerate(vals,1): wma.cell(r,c).value=v
# Add Methods + SI actions.
for vals in [
    ['CRITICAL','Methods thermal protocol','Methods still state 1000 trajectories/T with 10 burn + 10 measured cycles.','Replace with N=500 full-grid stationary series, second N=500 seed at T*=0.50-0.70, 60 burn + 100 measured cycles, trajectory-level inference and high-T dt checks.','C15-C19'],
    ['CRITICAL','SI S8','SI Tables S8a/S8b contain the retired short-window thermal boundaries.','Replace with the long-window stationary series, pooled high-T joint inference, stationarity/memory-loss and timestep-compatibility description.','C15-C19']
]:
    rr=wma.max_row+1
    for c in range(1,wma.max_column+1):
        src=wma.cell(rr-1,c); dst=wma.cell(rr,c)
        if src.has_style: dst._style=copy(src._style)
        dst.font=copy(src.font); dst.fill=copy(src.fill); dst.border=copy(src.border); dst.alignment=copy(src.alignment); dst.number_format=src.number_format
    for c,v in enumerate(vals,1): wma.cell(rr,c).value=v

# Keep sheet views and sensible widths unchanged; save.
wb.save(DST)

# Export current matrix sheet to CSV/JSON/MD for transparent audit.
wb2=load_workbook(DST, data_only=False)
ws2=wb2['Claim-Evidence Matrix']
rows=[]
for r in range(2,ws2.max_row+1):
    rows.append({headers[c-1]:ws2.cell(r,c).value for c in range(1,len(headers)+1)})
with open(OUT/'CLAIM_EVIDENCE_MATRIX_v2.csv','w',newline='',encoding='utf-8-sig') as f:
    w=csv.DictWriter(f,fieldnames=headers); w.writeheader(); w.writerows(rows)
with open(OUT/'CLAIM_EVIDENCE_MATRIX_v2.json','w',encoding='utf-8') as f:
    json.dump(rows,f,indent=2,ensure_ascii=False)
with open(OUT/'CLAIM_EVIDENCE_MATRIX_v2.md','w',encoding='utf-8') as f:
    f.write('# Janus MoSSe Claim-Evidence Matrix v2\n\n')
    f.write('Updated after RT-B01 thermal stationarity closure. Gate 11 Angeli provenance correction is also synchronized.\n\n')
    f.write('| ID | Category | Claim | Evidence | Scope |\n|---|---|---|---|---|\n')
    for row in rows:
        esc=lambda x: str(x or '').replace('|','\\|').replace('\n',' ')
        f.write(f"| {esc(row['ID'])} | {esc(row['Category'])} | {esc(row['Claim'])} | {esc(row['Evidence Status'])} | {esc(row['Scope'])} |\n")

# Markdown action register from updated sheet.
with open(OUT/'MANUSCRIPT_ACTION_REGISTER_v2.md','w',encoding='utf-8') as f:
    f.write('# Manuscript Action Register v2\n\n')
    for r in range(2,wma.max_row+1):
        vals=[wma.cell(r,c).value for c in range(1,wma.max_column+1)]
        f.write(f"- **{vals[0]} - {vals[1]}**: {vals[2]} Required: {vals[3]} ({vals[4]})\n")

print(DST)
