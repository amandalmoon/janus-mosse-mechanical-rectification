from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from copy import deepcopy
import re, shutil

src='/mnt/data/janus_n17_upstream_strengthening/Janus_MoSSe_UpstreamStrengthened_Manuscript_v5.docx'
out='/mnt/data/janus_acs_nano_venue/Janus_MoSSe_ACSNano_Structured_Manuscript_v6.docx'
shutil.copy2(src,out)
doc=Document(out)
body=doc._body._body

def element_text(el):
    return ''.join(t.text or '' for t in el.xpath('.//w:t')).strip()

def find_p_exact(text):
    for p in doc.paragraphs:
        if p.text.strip()==text:
            return p
    raise KeyError(text)

def delete_paragraph(p):
    parent=p._p.getparent(); parent.remove(p._p); p._p._p=None; p._element=None

# Remove working subtitle (not part of ACS title) and unheaded Introduction label.
for target in [
    'A DFT-anchored finite-contact model with prepared-state, symmetry, statistical, and in-plane relaxation falsification tests',
    '1. Introduction'
]:
    for p in list(doc.paragraphs):
        if p.text.strip()==target:
            delete_paragraph(p); break

# Keywords: 7 concise lowercase descriptors, preserving chemical formula via descriptive wording instead.
for p in doc.paragraphs:
    if p.text.strip().startswith('Keywords:'):
        p.text='Keywords: janus transition-metal dichalcogenides; mechanical rectification; depinning; finite-contact scaling; elastic relaxation; vector mode locking; rocking ratchet'
        break

# Strip numeric prefixes from manuscript section/subsection headings; pluralize Conclusions.
for p in doc.paragraphs:
    if p.style.name.startswith('Heading'):
        t=p.text.strip()
        t=re.sub(r'^\d+(?:\.\d+)*\.\s*', '', t)
        if t=='Model and computational methods': t='Methods'
        if t=='Conclusion': t='Conclusions'
        p.text=t

# Move the full Methods block from its current location to after Conclusions and before Data and Code Availability.
# Operate on body XML to preserve equations, figures, and tables.
children=list(body)
method_start=next(i for i,e in enumerate(children) if element_text(e)=='Methods')
results_start=next(i for i,e in enumerate(children) if element_text(e)=='Results')
methods_block=children[method_start:results_start]
for e in methods_block:
    body.remove(e)
# Recompute destination after removal.
children=list(body)
data_idx=next(i for i,e in enumerate(children) if element_text(e)=='Data and code availability' or element_text(e)=='Data and Code Availability')
for j,e in enumerate(methods_block):
    body.insert(data_idx+j,e)

# Normalize back-matter heading capitalization.
for p in doc.paragraphs:
    if p.text.strip().lower()=='data and code availability':
        p.text='Data and Code Availability'

# Insert official Supporting Information description before References.
ref_p=next(p for p in doc.paragraphs if p.text.strip()=='References')
si_p=doc.add_paragraph()
r=si_p.add_run('Supporting Information: '); r.bold=True
si_p.add_run('Numerical validation of prepared-state depinning, symmetry controls, size/shape and boundary sensitivity, deterministic and Floquet locking, thermal trajectory-cluster inference, and linear/nonlinear in-plane relaxation audits (PDF).')
si_p.paragraph_format.space_before = ref_p.paragraph_format.space_before
si_p.paragraph_format.space_after = ref_p.paragraph_format.space_after
# Move before References.
si_p._p.getparent().remove(si_p._p)
ref_p._p.addprevious(si_p._p)

# Keep title and abstract label readable; no author data are invented.
doc.save(out)
print(out)
