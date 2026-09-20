from docx import Document
from docx.oxml import OxmlElement
from pathlib import Path

src=Path('/mnt/data/janus_acs_nano_venue/Janus_MoSSe_ACSNano_Structured_Manuscript_v6.docx')
out=Path('/mnt/data/janus_acs_nano_venue/Janus_MoSSe_ACSNano_Structured_Manuscript_v7.docx')
doc=Document(src)

# Locate References heading; insert Acknowledgments immediately before it.
ref_para=None
for p in doc.paragraphs:
    if p.text.strip()=='References' and p.style.name.startswith('Heading'):
        ref_para=p
        break
if ref_para is None:
    raise RuntimeError('References heading not found')

# Avoid duplicate insertion.
if not any(p.text.strip()=='Acknowledgments' for p in doc.paragraphs):
    h=doc.add_paragraph(style='Heading 1')
    h.add_run('Acknowledgments')
    body=doc.add_paragraph(style='Body Text')
    body.add_run(
        'AI Assistance Disclosure. OpenAI ChatGPT was used to assist with research-workflow organization, '
        'code-assisted numerical auditing, literature-search organization, and manuscript drafting and editing. '
        'The authors independently reviewed and verified the scientific claims, calculations, citations, and final '
        'manuscript text and take responsibility for the submitted content.'
    )
    # Move newly created paragraphs before References, preserving their order.
    ref_para._p.addprevious(h._p)
    ref_para._p.addprevious(body._p)

doc.save(out)
print(out)
